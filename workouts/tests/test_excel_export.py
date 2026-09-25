"""Выгрузка истории в Excel: состав файла, формат строк, изоляция данных."""

import io
from datetime import datetime, timedelta

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.tests.factories import UserFactory
from workouts import excel
from workouts.models import Exercise, Sport
from workouts.tests.factories import (
    CardioPartFactory,
    ExerciseFactory,
    ExerciseNoteFactory,
    LocationFactory,
    RepsSetFactory,
    SportFactory,
    StrengthSetFactory,
    TimeSetFactory,
    WorkoutFactory,
)

pytestmark = pytest.mark.django_db


def moment(day, hour=18):
    """Момент тренировки в проектной таймзоне: даты в файле — локальные."""
    return timezone.make_aware(datetime(2026, 9, day, hour, 30))


def sheet_of(user):
    """Лист с историей, собранный напрямую — без HTTP."""
    book = excel.build_workbook(user)
    return book[excel.SHEET_TITLE]


def rows_of(sheet):
    """Строки листа словарями по ключам COLUMNS."""
    keys = [column.key for column in excel.COLUMNS]
    return [
        dict(zip(keys, values, strict=True))
        for values in sheet.iter_rows(min_row=2, values_only=True)
    ]


def test_export_requires_login(client):
    response = client.get(reverse("workout_export"))

    assert response.status_code == 302
    assert reverse("account_login") in response["Location"]


def test_export_returns_xlsx_with_dated_filename(client, user):
    client.force_login(user)

    response = client.get(reverse("workout_export"))
    content = b"".join(response.streaming_content)

    assert response.status_code == 200
    assert response["Content-Type"] == excel.CONTENT_TYPE
    assert "attachment" in response["Content-Disposition"]
    assert f"{timezone.localdate():%Y-%m-%d}" in response["Content-Disposition"]
    # Книга открывается, а не просто «какие-то байты приехали».
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == [excel.SHEET_TITLE, excel.HELP_SHEET_TITLE]


def test_header_matches_columns(user):
    sheet = sheet_of(user)

    titles = [cell.value for cell in sheet[1]]

    assert titles == [column.title for column in excel.COLUMNS]


def test_empty_history_gives_header_only(user):
    sheet = sheet_of(user)

    assert sheet.max_row == 1


def test_strength_workout_gives_row_per_set(user):
    workout = WorkoutFactory(user=user, started_at=moment(4), duration_min=89)
    bench = ExerciseFactory(name="Жим лёжа")
    StrengthSetFactory(workout=workout, exercise=bench, set_number=1, weight_kg=80, reps=8)
    StrengthSetFactory(workout=workout, exercise=bench, set_number=2, weight_kg=82.5, reps=6)

    rows = rows_of(sheet_of(user))

    assert len(rows) == 2
    assert [row["set_number"] for row in rows] == [1, 2]
    assert [float(row["weight"]) for row in rows] == [80.0, 82.5]
    assert [row["reps"] for row in rows] == [8, 6]
    assert {row["exercise"] for row in rows} == {"Жим лёжа"}
    # Длительность повторяется в каждой строке: человек читает лист построчно.
    assert [row["duration"] for row in rows] == [89, 89]
    assert [row["date"] for row in rows] == [moment(4).date(), moment(4).date()]


@pytest.mark.parametrize(
    ("factory", "measurement", "filled", "empty"),
    [
        (StrengthSetFactory, Exercise.Measurement.WEIGHT_REPS, ("weight", "reps"), ("hold",)),
        (RepsSetFactory, Exercise.Measurement.REPS, ("reps",), ("weight", "hold")),
        (TimeSetFactory, Exercise.Measurement.TIME, ("hold",), ("weight", "reps")),
    ],
)
def test_columns_match_measurement(user, factory, measurement, filled, empty):
    """Заполнены ровно те колонки, которые осмысленны для единицы упражнения."""
    workout = WorkoutFactory(user=user, started_at=moment(4))
    exercise = ExerciseFactory(measurement=measurement)
    factory(workout=workout, exercise=exercise)

    row = rows_of(sheet_of(user))[0]

    assert all(row[key] not in (None, "") for key in filled)
    assert all(row[key] in (None, "") for key in empty)


def test_time_weight_set_keeps_both_values(user):
    workout = WorkoutFactory(user=user, started_at=moment(4))
    exercise = ExerciseFactory(measurement=Exercise.Measurement.TIME_WEIGHT)
    TimeSetFactory(
        workout=workout,
        exercise=exercise,
        measurement=Exercise.Measurement.TIME_WEIGHT,
        duration_sec=90,
        weight_kg=20,
        reps=0,
    )

    row = rows_of(sheet_of(user))[0]

    # Удержание — «1:30», тем же форматом, каким его видно в приложении.
    assert row["hold"] == "1:30"
    assert float(row["weight"]) == 20.0
    assert row["reps"] in (None, "")


def test_cardio_workout_gives_single_row(user):
    sport = SportFactory(name="Велосипед", category=Sport.Category.CARDIO)
    workout = WorkoutFactory(user=user, sport=sport, started_at=moment(3), duration_min=60)
    CardioPartFactory(workout=workout, distance_km=24.5, avg_heart_rate=142)

    rows = rows_of(sheet_of(user))

    assert len(rows) == 1
    row = rows[0]
    assert row["sport"] == "Велосипед"
    assert row["exercise"] in (None, "")
    assert float(row["distance"]) == 24.5
    assert row["pulse"] == 142
    assert row["duration"] == 60


def test_notes_are_written_once(user):
    workout = WorkoutFactory(user=user, started_at=moment(4), note="Тяжело зашло")
    bench = ExerciseFactory(name="Жим лёжа")
    StrengthSetFactory(workout=workout, exercise=bench, set_number=1)
    StrengthSetFactory(workout=workout, exercise=bench, set_number=2)
    ExerciseNoteFactory(workout=workout, exercise=bench, text="Узкий хват")

    rows = rows_of(sheet_of(user))

    assert [row["workout_note"] for row in rows] == ["Тяжело зашло", None]
    assert [row["exercise_note"] for row in rows] == ["Узкий хват", None]


def test_location_is_written(user):
    place = LocationFactory(owner=user, name="СпортЛайф")
    workout = WorkoutFactory(user=user, started_at=moment(4), location=place)
    StrengthSetFactory(workout=workout)

    assert rows_of(sheet_of(user))[0]["location"] == "СпортЛайф"


def test_formula_like_note_stays_text(user):
    """«=2+2» в заметке — это заметка, а не формула, которую Excel посчитает."""
    workout = WorkoutFactory(user=user, started_at=moment(4), note="=2+2")
    StrengthSetFactory(workout=workout)

    sheet = sheet_of(user)
    note_column = [column.key for column in excel.COLUMNS].index("workout_note") + 1
    cell = sheet.cell(row=2, column=note_column)

    assert cell.value == "=2+2"
    assert cell.data_type == "s"


def test_drafts_and_live_workouts_are_not_exported(user):
    WorkoutFactory(user=user, started_at=None, duration_min=None)
    WorkoutFactory(user=user, started_at=moment(4), duration_min=None)
    recorded = WorkoutFactory(user=user, started_at=moment(3))
    StrengthSetFactory(workout=recorded)

    rows = rows_of(sheet_of(user))

    assert len(rows) == 1
    assert rows[0]["date"] == moment(3).date()


def test_rows_are_ordered_by_date(user):
    for day in (5, 3, 4):
        workout = WorkoutFactory(user=user, started_at=moment(day))
        StrengthSetFactory(workout=workout)

    rows = rows_of(sheet_of(user))

    assert [row["date"].day for row in rows] == [3, 4, 5]


def test_exercise_order_follows_the_workout(user):
    """Порядок упражнений в файле — тот же, что на экране итога."""
    workout = WorkoutFactory(user=user, started_at=moment(4))
    squat = ExerciseFactory(name="Присед")
    bench = ExerciseFactory(name="Жим лёжа")
    # Жим добавлен раньше, но выполнен позже — в итоге он идёт вторым.
    StrengthSetFactory(
        workout=workout, exercise=bench, set_number=1, done_at=moment(4) + timedelta(minutes=40)
    )
    StrengthSetFactory(
        workout=workout, exercise=squat, set_number=1, done_at=moment(4) + timedelta(minutes=5)
    )

    rows = rows_of(sheet_of(user))

    assert [row["exercise"] for row in rows] == ["Присед", "Жим лёжа"]


def test_export_does_not_include_other_users_workouts(user):
    stranger = UserFactory()
    mine = WorkoutFactory(user=user, started_at=moment(4))
    StrengthSetFactory(workout=mine, exercise=ExerciseFactory(name="Моё упражнение"))
    theirs = WorkoutFactory(user=stranger, started_at=moment(4))
    StrengthSetFactory(workout=theirs, exercise=ExerciseFactory(name="Чужое упражнение"))

    rows = rows_of(sheet_of(user))

    assert [row["exercise"] for row in rows] == ["Моё упражнение"]


def test_help_sheet_is_present(user):
    book = excel.build_workbook(user)

    lines = [row[0] for row in book[excel.HELP_SHEET_TITLE].iter_rows(values_only=True)]

    assert lines[0] == "Как заполнять таблицу"
    assert any("Кардио" in (line or "") for line in lines)
