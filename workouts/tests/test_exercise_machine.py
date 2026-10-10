"""Тренажёр упражнения: производитель и модель — у пары «упражнение × место».

Изоляция: место — только своё, упражнение — только видимое, тренажёры у
каждого свои, даже у общего упражнения.
"""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from workouts import exercise_excel
from workouts.models import ExerciseMachine, Location, MachineBrand, MachineModel
from workouts.tests.factories import (
    ExerciseFactory,
    ExerciseMachineFactory,
    LocationFactory,
    MachineBrandFactory,
    MachineModelFactory,
    StrengthSetFactory,
    WorkoutFactory,
)

pytestmark = pytest.mark.django_db


@pytest.fixture
def press(db):
    return ExerciseFactory(name="Жим ногами", owner=None, equipment="Тренажёр")


@pytest.fixture
def gym(user):
    return LocationFactory(owner=user, name="СпортЛайф")


def url(exercise, location):
    return reverse("exercise_machine", args=[exercise.pk, location.pk])


def brand(name, owner=None):
    return MachineBrandFactory(name=name, owner=owner)


def post(client, exercise, location, **data):
    return client.post(url(exercise, location), data)


def machine_of(user, exercise, location):
    return ExerciseMachine.objects.filter(user=user, exercise=exercise, location=location).first()


# ---------- Окно «Тренажёр» ----------


def test_modal_starts_with_brands_and_opens_current_brand(client, user, other_user, press, gym):
    brand("Technogym")
    brand("Свой бренд", owner=user)
    brand("Чужой бренд", owner=other_user)
    client.force_login(user)

    html = client.get(url(press, gym)).content.decode()

    assert "Technogym" in html
    assert "Свой бренд" in html
    assert "Чужой бренд" not in html

    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    html = client.get(url(press, gym)).content.decode()
    assert "Selection 900" in html  # сразу шаг моделей текущего производителя
    assert "Другой производитель" in html


def test_pick_shared_brand_and_model(client, user, press, gym):
    model = MachineModelFactory(brand=brand("Technogym"), name="Pure Leg Press")
    client.force_login(user)

    response = post(client, press, gym, action="pick", brand=model.brand_id, model=model.pk)

    assert response.headers["HX-Refresh"] == "true"
    machine = machine_of(user, press, gym)
    assert (machine.brand, machine.model) == (model.brand, model)


def test_pick_brand_without_model(client, user, press, gym):
    technogym = brand("Technogym")
    client.force_login(user)

    post(client, press, gym, action="pick", brand=technogym.pk, model="")

    machine = machine_of(user, press, gym)
    assert machine.brand == technogym
    assert machine.model is None
    assert machine.label == "Technogym"


def test_new_brand_by_name_is_own_and_reuses_existing(client, user, press, gym):
    technogym = brand("Technogym")
    client.force_login(user)

    html = post(client, press, gym, action="brand_new", **{"brand-name": " TECHNOGYM "})
    assert MachineBrand.objects.count() == 1  # совпало имя — выбран общий
    assert "Модель не знаю" in html.content.decode()

    post(client, press, gym, action="brand_new", **{"brand-name": "Kettler"})
    kettler = MachineBrand.objects.get(name="Kettler")
    assert kettler.owner == user
    assert technogym.is_global


def test_new_model_is_own_and_saved(client, user, press, gym):
    technogym = brand("Technogym")
    client.force_login(user)

    response = post(
        client,
        press,
        gym,
        action="model_new",
        brand=technogym.pk,
        **{"model-name": "Selection  700"},
    )

    assert response.headers["HX-Refresh"] == "true"
    machine = machine_of(user, press, gym)
    assert machine.model.name == "Selection 700"
    assert machine.model.owner == user
    assert machine.model.brand == technogym


def test_regular_user_cannot_create_shared(client, user, press, gym):
    client.force_login(user)

    post(
        client, press, gym, action="brand_new", **{"brand-name": "Matrix", "brand-scope": "global"}
    )

    assert MachineBrand.objects.get(name="Matrix").owner == user


def test_admin_creates_shared_brand_and_model(client, user, press, gym):
    user.is_staff = True
    user.save(update_fields=["is_staff"])
    client.force_login(user)

    post(
        client, press, gym, action="brand_new", **{"brand-name": "Matrix", "brand-scope": "global"}
    )
    matrix = MachineBrand.objects.get(name="Matrix")
    assert matrix.is_global

    post(
        client,
        press,
        gym,
        action="model_new",
        brand=matrix.pk,
        **{"model-name": "Ultra", "model-scope": "global"},
    )
    assert MachineModel.objects.get(name="Ultra").is_global


def test_admin_shared_model_needs_shared_brand(client, user, press, gym):
    user.is_staff = True
    user.save(update_fields=["is_staff"])
    own = brand("Свой", owner=user)
    client.force_login(user)

    response = post(
        client,
        press,
        gym,
        action="model_new",
        brand=own.pk,
        **{"model-name": "Ultra", "model-scope": "global"},
    )

    assert "HX-Refresh" not in response.headers
    assert "только у общего производителя" in response.content.decode()
    assert not MachineModel.objects.exists()


def test_remove_machine(client, user, press, gym):
    ExerciseMachineFactory(user=user, exercise=press, location=gym)
    client.force_login(user)

    post(client, press, gym, action="remove")

    assert not ExerciseMachine.objects.exists()


def test_someone_elses_brand_or_model_is_404(client, user, other_user, press, gym):
    theirs = brand("Чужой", owner=other_user)
    their_model = MachineModelFactory(brand=brand("Technogym"), name="Чужая", owner=other_user)
    client.force_login(user)

    assert client.get(url(press, gym), {"brand": theirs.pk}).status_code == 404
    assert post(client, press, gym, action="pick", brand=theirs.pk).status_code == 404
    response = post(
        client, press, gym, action="pick", brand=their_model.brand_id, model=their_model.pk
    )
    assert response.status_code == 404
    assert not ExerciseMachine.objects.exists()


def test_model_of_another_brand_is_404(client, user, press, gym):
    model = MachineModelFactory(brand=brand("Technogym"), name="Pure")
    other = brand("Matrix")
    client.force_login(user)

    assert (
        post(client, press, gym, action="pick", brand=other.pk, model=model.pk).status_code == 404
    )


def test_someone_elses_location_is_404(client, user, other_user, press):
    theirs = LocationFactory(owner=other_user, name="Чужой зал")
    client.force_login(user)

    assert client.get(url(press, theirs)).status_code == 404
    technogym = brand("Technogym")
    assert post(client, press, theirs, action="pick", brand=technogym.pk).status_code == 404
    assert not ExerciseMachine.objects.exists()


def test_someone_elses_personal_exercise_is_404(client, user, other_user, gym):
    theirs = ExerciseFactory(name="Секретное", owner=other_user)
    client.force_login(user)

    technogym = brand("Technogym")
    assert post(client, theirs, gym, action="pick", brand=technogym.pk).status_code == 404
    assert not ExerciseMachine.objects.exists()


def test_guest_is_sent_to_login(client, press, gym):
    response = client.get(url(press, gym))

    assert response.status_code == 302
    assert response.url.startswith(reverse("account_login"))


# ---------- Страница упражнения ----------


def detail(client, exercise):
    return client.get(reverse("exercise_detail", args=[exercise.pk])).content.decode()


def test_exercise_page_lists_own_places_with_machines(client, user, other_user, press, gym):
    LocationFactory(owner=user, name="Дом")
    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    ExerciseMachineFactory(
        user=other_user, exercise=press, brand=brand("Чужой бренд", owner=other_user)
    )
    client.force_login(user)

    html = detail(client, press)

    assert "Hammer Strength · Selection 900" in html
    assert "Дом" in html
    assert "тренажёр не указан" in html
    assert "Чужой бренд" not in html


def test_shared_exercise_has_own_machine_per_user(client, user, other_user, press, gym):
    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    ExerciseMachineFactory(
        user=other_user,
        exercise=press,
        model=MachineModelFactory(brand=brand("Technogym"), name="Pure"),
        brand=brand("Technogym"),
    )

    client.force_login(other_user)
    html = detail(client, press)

    assert "Technogym · Pure" in html
    assert "Hammer Strength" not in html


def test_section_without_places_points_to_my_locations(client, user, press):
    client.force_login(user)

    html = detail(client, press)

    assert reverse("my_locations") in html


def test_exercise_history_shows_machine_of_that_place(client, user, press, gym):
    StrengthSetFactory(workout=WorkoutFactory(user=user, location=gym), exercise=press)
    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    client.force_login(user)

    html = detail(client, press)

    assert "app-machine-line" in html  # строка тренажёра у записи истории


# ---------- Живой режим, итог, правка ----------


def live_workout(user, location, exercise):
    workout = WorkoutFactory(user=user, location=location, duration_min=None)
    StrengthSetFactory(workout=workout, exercise=exercise, done=False)
    return workout


def test_live_screen_shows_machine_of_workout_place(client, user, press, gym):
    other_place = LocationFactory(owner=user, name="Дом")
    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    ExerciseMachineFactory(user=user, exercise=press, location=other_place, brand=brand("Kettler"))
    workout = live_workout(user, gym, press)
    client.force_login(user)

    html = client.get(reverse("workout_live", args=[workout.pk])).content.decode()

    assert "Hammer Strength · Selection 900" in html
    assert "Kettler" not in html


def test_live_screen_offers_machine_for_machine_exercises(client, user, press, gym):
    workout = live_workout(user, gym, press)
    client.force_login(user)

    html = client.get(reverse("workout_live", args=[workout.pk])).content.decode()

    assert "Указать тренажёр" in html
    assert url(press, gym) in html


def test_live_screen_offers_maker_for_barbell_and_dumbbells(client, user, gym):
    """У блинов и гантелей производитель тоже есть — кнопка «Производитель»."""
    squat = ExerciseFactory(name="Присед", owner=None, equipment="Штанга")
    workout = live_workout(user, gym, squat)
    client.force_login(user)

    html = client.get(reverse("workout_live", args=[workout.pk])).content.decode()

    assert "Указать производителя" in html
    assert "Указать тренажёр" not in html
    assert url(squat, gym) in html


@pytest.mark.parametrize("equipment", ["Своё тело", ""])
def test_live_screen_does_not_offer_maker_for_bodyweight(client, user, gym, equipment):
    pullup = ExerciseFactory(name="Подтягивания", owner=None, equipment=equipment)
    workout = live_workout(user, gym, pullup)
    client.force_login(user)

    html = client.get(reverse("workout_live", args=[workout.pk])).content.decode()

    assert "Указать производителя" not in html
    assert "Указать тренажёр" not in html


def test_dumbbell_page_and_modal_say_maker_not_machine(client, user, gym):
    dumbbells = ExerciseFactory(name="Жим гантелей", owner=None, equipment="Гантели")
    client.force_login(user)

    page = client.get(reverse("exercise_detail", args=[dumbbells.pk])).content.decode()
    modal = client.get(url(dumbbells, gym)).content.decode()

    assert "производитель не указан" in page
    assert "тренажёр не указан" not in page
    assert "Производитель и модель для «Жим гантелей»" in modal


def test_workout_without_place_has_no_machine(client, user, press):
    workout = live_workout(user, None, press)
    client.force_login(user)

    html = client.get(reverse("workout_live", args=[workout.pk])).content.decode()

    assert "Указать тренажёр" not in html
    assert "exercise_machine" not in html


def test_summary_shows_machine(client, user, press, gym):
    workout = WorkoutFactory(user=user, location=gym)
    StrengthSetFactory(workout=workout, exercise=press)
    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    client.force_login(user)

    html = client.get(reverse("workout_summary", args=[workout.pk])).content.decode()

    assert "Hammer Strength · Selection 900" in html


def test_deleting_place_removes_its_machines(client, user, press):
    place = LocationFactory(owner=user, name="Старый зал")
    ExerciseMachineFactory(user=user, exercise=press, location=place)
    client.force_login(user)

    client.post(reverse("location_delete", args=[place.pk]))

    assert not Location.objects.filter(pk=place.pk).exists()
    assert not ExerciseMachine.objects.exists()


def screen_queries(client, user, count):
    """Сколько запросов стоят живой экран, итог и правка при count упражнений
    с тренажёрами — у отдельного пользователя, чтобы наборы не смешивались."""
    gym = LocationFactory(owner=user, name="Зал")
    live = WorkoutFactory(user=user, location=gym, duration_min=None)
    done = WorkoutFactory(user=user, location=gym)
    for number in range(count):
        exercise = ExerciseFactory(name=f"Тренажёр {user.pk}-{number}", equipment="Тренажёр")
        StrengthSetFactory(workout=live, exercise=exercise, done=False)
        StrengthSetFactory(workout=done, exercise=exercise)
        ExerciseMachineFactory(user=user, exercise=exercise, location=gym)
    client.force_login(user)
    counts = []
    for name, workout in [
        ("workout_live", live),
        ("workout_summary", done),
        ("workout_correct", done),
    ]:
        with CaptureQueriesContext(connection) as queries:
            assert client.get(reverse(name, args=[workout.pk])).status_code == 200
        counts.append(len(queries))
    return counts


def test_screens_budgets_do_not_grow_with_machines(client, user, other_user):
    """Тренажёр приходит подзапросом в выборке подходов — экранам он бесплатен:
    одно упражнение с тренажёром и шесть стоят одинаково."""
    assert screen_queries(client, user, 1) == screen_queries(client, other_user, 6)


# ---------- Excel справочника ----------


def round_trip(client, user, edit=None):
    buffer = io.BytesIO()
    book = exercise_excel.build_workbook(user)
    if edit:
        edit(book[exercise_excel.MACHINE_SHEET_TITLE])
    book.save(buffer)
    upload = SimpleUploadedFile("справочник.xlsx", buffer.getvalue())
    client.force_login(user)
    response = client.post(reverse("exercise_import"), {"exercises-file": upload})
    return response.context["report"]


def test_export_has_machine_sheet(user, press, gym):
    ExerciseMachineFactory(user=user, exercise=press, location=gym, brand=brand("Hammer Strength"))
    buffer = io.BytesIO()
    exercise_excel.build_workbook(user).save(buffer)

    sheet = load_workbook(buffer)[exercise_excel.MACHINE_SHEET_TITLE]
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == ("ID упражнения", "Упражнение", "Место", "Производитель", "Модель")
    assert rows[1] == (press.pk, "Жим ногами", "СпортЛайф", "Hammer Strength", "Selection 900")


def test_untouched_machine_sheet_writes_nothing(client, user, press, gym):
    machine = ExerciseMachineFactory(user=user, exercise=press, location=gym)

    report = round_trip(client, user)

    assert report.machines == 0
    assert not report.errors
    assert ExerciseMachine.objects.get() == machine


def test_machine_sheet_edits_are_applied(client, user, press, gym):
    ExerciseMachineFactory(user=user, exercise=press, location=gym)
    LocationFactory(owner=user, name="Дом")

    def edit(sheet):
        sheet["E2"] = "Selection 700"
        sheet.append([None, "Жим ногами", "дом", "Kettler", ""])

    report = round_trip(client, user, edit)

    assert report.machines == 2
    assert machine_of(user, press, gym).model.name == "Selection 700"
    assert ExerciseMachine.objects.get(location__name="Дом").brand.name == "Kettler"


def test_machine_sheet_empty_fields_remove_machine(client, user, press, gym):
    ExerciseMachineFactory(user=user, exercise=press, location=gym)

    def edit(sheet):
        sheet["D2"] = None
        sheet["E2"] = None

    round_trip(client, user, edit)

    assert not ExerciseMachine.objects.exists()


def test_machine_sheet_unknown_place_is_a_row_error(client, user, other_user, press):
    LocationFactory(owner=other_user, name="Чужой зал")

    def edit(sheet):
        sheet.append([press.pk, "Жим ногами", "Чужой зал", "Technogym", ""])

    report = round_trip(client, user, edit)

    assert "Чужой зал" in report.errors[0]
    assert "листа «Тренажёры»" in report.errors[0]
    assert not ExerciseMachine.objects.exists()
    assert not Location.objects.filter(owner=user).exists()


def test_file_without_machine_sheet_keeps_machines(client, user, press, gym):
    ExerciseMachineFactory(user=user, exercise=press, location=gym)
    buffer = io.BytesIO()
    book = exercise_excel.build_workbook(user)
    del book[exercise_excel.MACHINE_SHEET_TITLE]
    book.save(buffer)
    client.force_login(user)

    client.post(
        reverse("exercise_import"),
        {"exercises-file": SimpleUploadedFile("старый.xlsx", buffer.getvalue())},
    )

    assert ExerciseMachine.objects.count() == 1
