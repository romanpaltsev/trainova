"""Круги: суперсет, трисет, круг из N упражнений.

Круг — упражнения тренировки с одним номером StrengthSet.circuit. Правило
порядка одно на три потребителя (models.order_exercises): итог, номер
упражнения на его странице и подпись тренировки по группам мышц. Подходы
круга идут по раундам A1 → B1 → A2 → B2, отдых — только после раунда.
"""

import io
from datetime import timedelta

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from workouts import excel, excel_import, services, stats, trash
from workouts.models import DeletedWorkout, StrengthSet, Workout
from workouts.tests.budgets import SIDEBAR_QUERIES
from workouts.tests.factories import ExerciseFactory, StrengthSetFactory, WorkoutFactory

pytestmark = pytest.mark.django_db


def make(workout, exercise, count=2, *, circuit=None, done=None):
    """Подходы упражнения; done — сколько первых уже выполнено."""
    done = workout.is_finished if done is None else done
    rows = []
    for number in range(1, count + 1):
        is_done = done is True or (isinstance(done, int) and number <= done)
        rows.append(
            StrengthSetFactory(
                workout=workout,
                exercise=exercise,
                set_number=number,
                circuit=circuit,
                done=bool(is_done),
            )
        )
    return rows


def names(groups):
    return [group["exercise"].name for group in groups]


# ---------- Одно правило порядка ----------


def test_circuit_members_go_together_and_three_consumers_agree(user):
    """Добавлены A, B, C, D; в круге A и C. Порядок — [A, C], B, D: круг
    встаёт на место своего первого члена. Итог, страница упражнения и подпись
    обязаны совпасть — в том числе с членом круга без группы мышц."""
    a = ExerciseFactory(name="A", muscle_group="Грудь")
    b = ExerciseFactory(name="B", muscle_group="Ноги")
    c = ExerciseFactory(name="C", muscle_group="")
    d = ExerciseFactory(name="D", muscle_group="Спина")
    workout = WorkoutFactory(user=user)
    for exercise in (a, b, c, d):
        make(workout, exercise, circuit=1 if exercise in (a, c) else None)

    groups = services.exercise_groups(workout)
    positions = {
        exercise.name: stats.exercise_positions(user, [workout.pk], exercise)[workout.pk]
        for exercise in (a, b, c, d)
    }

    assert names(groups) == ["A", "C", "B", "D"]
    assert positions == {group["exercise"].name: group["position"] for group in groups}
    assert stats.muscle_groups_by_workout(user, [workout.pk]) == {workout.pk: "Грудь · Ноги +1"}
    assert [group["circuit_no"] for group in groups] == [1, 1, None, None]


def test_order_follows_facts_inside_and_between_blocks(user):
    """Выполненное раньше — раньше: одиночный B сделан первым и встаёт перед
    кругом; внутри круга C сделан раньше A."""
    a, b, c = (ExerciseFactory(name=name) for name in "ABC")
    workout = WorkoutFactory(user=user)
    now = timezone.now()
    for exercise, minutes in ((a, 10), (b, 1), (c, 5)):
        row = make(workout, exercise, count=1, circuit=1 if exercise != b else None)[0]
        StrengthSet.objects.filter(pk=row.pk).update(done_at=now + timedelta(minutes=minutes))

    assert names(services.exercise_groups(workout)) == ["B", "C", "A"]


def test_lone_circuit_number_is_not_a_circuit(user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user)
    make(workout, a, circuit=7)
    make(workout, b)

    groups = services.exercise_groups(workout)

    assert [group["circuit"] for group in groups] == [None, None]


def test_workout_without_circuits_keeps_old_order(user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user)
    make(workout, b)
    make(workout, a)

    assert names(services.exercise_groups(workout)) == ["B", "A"]


# ---------- Связать и разобрать ----------


@pytest.fixture
def live(user):
    workout = WorkoutFactory(user=user, duration_min=None)
    exercises = [ExerciseFactory(name=name) for name in "ABCD"]
    for exercise in exercises:
        make(workout, exercise, done=False)
    return workout, exercises


def circuits_of(workout):
    return {group["exercise"].name: group["circuit"] for group in services.exercise_groups(workout)}


def post_circuit(client, workout, **payload):
    return client.post(reverse("live_circuits", args=[workout.pk]), payload)


def test_link_creates_joins_and_merges_circuits(client, user, live):
    workout, (a, b, c, d) = live
    client.force_login(user)

    post_circuit(client, workout, action="link", exercise=a.pk, next=b.pk)
    first = circuits_of(workout)
    assert first["A"] == first["B"] and first["A"] is not None and first["C"] is None

    post_circuit(client, workout, action="link", exercise=c.pk, next=d.pk)
    post_circuit(client, workout, action="link", exercise=b.pk, next=c.pk)

    assert len(set(circuits_of(workout).values())) == 1


def test_stale_or_non_adjacent_pair_changes_nothing(client, user, live):
    workout, (a, _b, c, _d) = live
    client.force_login(user)

    response = post_circuit(client, workout, action="link", exercise=a.pk, next=c.pk)

    assert response.status_code == 200
    assert set(circuits_of(workout).values()) == {None}


def test_unlink_dissolves_circuit_of_one(client, user, live):
    workout, (a, b, c, _d) = live
    client.force_login(user)
    services.link_with_next(workout, a.pk, b.pk)
    services.link_with_next(workout, b.pk, c.pk)

    post_circuit(client, workout, action="unlink", exercise=b.pk)
    after_one = circuits_of(workout)
    post_circuit(client, workout, action="unlink", exercise=a.pk)

    assert after_one["B"] is None and after_one["A"] == after_one["C"] is not None
    assert set(circuits_of(workout).values()) == {None}


def test_modal_answers_with_region_out_of_band(client, user, live):
    workout, (a, b, _c, _d) = live
    client.force_login(user)

    content = post_circuit(
        client, workout, action="link", exercise=a.pk, next=b.pk
    ).content.decode()

    assert "Суперсет" in content
    assert 'id="exercises" hx-swap-oob="true"' in content


def test_exercise_from_another_workout_is_404(client, user, live, other_user):
    workout, (a, _b, _c, _d) = live
    stranger = ExerciseFactory(name="Чужое")
    make(WorkoutFactory(user=other_user), stranger)
    client.force_login(user)

    response = post_circuit(client, workout, action="link", exercise=a.pk, next=stranger.pk)

    assert response.status_code == 404


def test_link_tile_appears_from_the_second_exercise(client, user):
    workout = WorkoutFactory(user=user, duration_min=None)
    make(workout, ExerciseFactory(name="A"), done=False)
    client.force_login(user)
    url = reverse("workout_live", args=[workout.pk])
    tile = reverse("live_circuits", args=[workout.pk])

    assert tile not in client.get(url).content.decode()
    make(workout, ExerciseFactory(name="B"), done=False)
    content = client.get(url).content.decode()
    assert tile in content
    assert "app-circuit" not in content  # кругов нет — обёрток тоже


def test_deleting_last_set_of_member_dissolves_circuit(client, user, live):
    workout, (a, b, _c, _d) = live
    services.link_with_next(workout, a.pk, b.pk)
    client.force_login(user)

    for row in workout.sets.filter(exercise=b):
        client.post(reverse("set_delete", args=[row.pk]))

    assert set(workout.sets.values_list("circuit", flat=True)) == {None}


# ---------- Раунды и отдых ----------


def done(client, workout, exercise, number):
    row = workout.sets.get(exercise=exercise, set_number=number)
    return client.post(reverse("set_done", args=[row.pk])).content.decode()


def test_sets_go_by_rounds_and_rest_only_after_a_round(client, user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user, duration_min=None)
    make(workout, a, done=False, circuit=1)
    make(workout, b, done=False, circuit=1)
    client.force_login(user)

    after_a1 = done(client, workout, a, 1)
    workout.refresh_from_db()
    assert workout.current_exercise_id == b.pk
    assert "data-stop" in after_a1 and "data-autostart" not in after_a1

    after_b1 = done(client, workout, b, 1)
    workout.refresh_from_db()
    assert workout.current_exercise_id == a.pk
    assert "data-autostart" in after_b1

    done(client, workout, a, 2)
    after_b2 = done(client, workout, b, 2)
    workout.refresh_from_db()
    assert workout.current_exercise_id is None
    assert "data-autostart" in after_b2


def test_manual_choice_inside_circuit_lasts_one_set(client, user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user, duration_min=None)
    make(workout, a, done=False, circuit=1)
    make(workout, b, done=False, circuit=1)
    client.force_login(user)
    client.post(reverse("live_exercise_select", args=[workout.pk]), {"exercise": b.pk})

    done(client, workout, b, 1)

    workout.refresh_from_db()
    assert workout.current_exercise_id == a.pk


def test_set_outside_circuit_restarts_rest_as_before(client, user):
    single = ExerciseFactory(name="Одиночное")
    workout = WorkoutFactory(user=user, duration_min=None)
    make(workout, single, done=False)
    client.force_login(user)

    content = done(client, workout, single, 1)

    assert "data-autostart" in content and "data-stop" not in content


# ---------- «+ Круг», «+ в круг», «Повторить» ----------


def test_round_adds_a_set_to_every_member(client, user, live):
    workout, (a, b, _c, _d) = live
    services.link_with_next(workout, a.pk, b.pk)
    number = workout.sets.filter(exercise=a).first().circuit
    client.force_login(user)

    client.post(reverse("live_round_add", args=[workout.pk]), {"circuit": number})

    assert workout.sets.filter(exercise=a).count() == 3
    assert workout.sets.filter(exercise=b).count() == 3
    assert set(workout.sets.filter(exercise__in=[a, b]).values_list("circuit", flat=True)) == {
        number
    }


def test_round_on_finished_workout_is_done_without_timestamp(client, user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user)
    make(workout, a, circuit=1)
    make(workout, b, circuit=1)
    client.force_login(user)

    client.post(reverse("live_round_add", args=[workout.pk]), {"circuit": 1})

    added = workout.sets.filter(set_number=3)
    assert added.count() == 2
    assert all(row.done and row.done_at is None for row in added)


def test_unknown_round_circuit_is_404(client, user, live):
    workout, _exercises = live
    client.force_login(user)

    response = client.post(reverse("live_round_add", args=[workout.pk]), {"circuit": 5})

    assert response.status_code == 404


def test_exercise_added_into_circuit(client, user, live):
    workout, (a, b, _c, _d) = live
    services.link_with_next(workout, a.pk, b.pk)
    number = workout.sets.filter(exercise=a).first().circuit
    newcomer = ExerciseFactory(name="Новое")
    client.force_login(user)

    picker = client.get(reverse("live_exercises", args=[workout.pk]), {"circuit": number})
    client.post(
        reverse("live_exercises", args=[workout.pk]), {"exercise": newcomer.pk, "circuit": number}
    )

    assert "Упражнение в круг" in picker.content.decode()
    assert set(workout.sets.filter(exercise=newcomer).values_list("circuit", flat=True)) == {number}


def test_stale_circuit_param_adds_exercise_alone(client, user, live):
    workout, _exercises = live
    newcomer = ExerciseFactory(name="Новое")
    client.force_login(user)

    client.post(
        reverse("live_exercises", args=[workout.pk]), {"exercise": newcomer.pk, "circuit": 9}
    )

    assert set(workout.sets.filter(exercise=newcomer).values_list("circuit", flat=True)) == {None}


def test_repeat_keeps_circuits(client, user):
    a, b, c = (ExerciseFactory(name=name) for name in "ABC")
    source = WorkoutFactory(user=user)
    make(source, a, circuit=4)
    make(source, b, circuit=4)
    make(source, c)
    client.force_login(user)

    client.post(reverse("workout_repeat", args=[source.pk]))

    repeated = Workout.objects.filter(user=user).live().get()
    assert circuits_of(repeated) == {"A": 1, "B": 1, "C": None}


# ---------- Итог, правка, корзина, Excel ----------


def test_summary_and_correct_screen_show_the_circuit(client, user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user)
    make(workout, a, circuit=1)
    make(workout, b, circuit=1)
    client.force_login(user)

    summary = client.get(reverse("workout_summary", args=[workout.pk])).content.decode()
    correct = client.get(reverse("workout_correct", args=[workout.pk])).content.decode()

    assert 'class="app-circuit" aria-label="Суперсет"' in summary
    assert "2 круга" in summary
    assert reverse("live_round_add", args=[workout.pk]) in correct


def test_circuit_survives_trash_and_restore(user):
    a, b = ExerciseFactory(name="A"), ExerciseFactory(name="B")
    workout = WorkoutFactory(user=user)
    make(workout, a, circuit=3)
    make(workout, b, circuit=3)
    entry = trash.move_to_trash(workout, title="Грудь", subtitle="")

    trash.restore(DeletedWorkout.objects.get(pk=entry.pk).pk, user)

    assert set(Workout.objects.get(pk=workout.pk).sets.values_list("circuit", flat=True)) == {3}


def test_excel_round_trip_keeps_circuits(user, other_user):
    a, b, c = (ExerciseFactory(name=name) for name in "ABC")
    workout = WorkoutFactory(user=user)
    make(workout, a, circuit=5)
    make(workout, b, circuit=5)
    make(workout, c)
    buffer = io.BytesIO()
    excel.build_workbook(user).save(buffer)

    sheet = load_workbook(io.BytesIO(buffer.getvalue()))[excel.SHEET_TITLE]
    header = [cell.value for cell in sheet[1]]
    circuits = {
        row[header.index("Упражнение")]: row[header.index("Круг")]
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    assert circuits == {"A": 1, "B": 1, "C": None}

    excel_import.import_workbook(other_user, io.BytesIO(buffer.getvalue()))
    imported = Workout.objects.get(user=other_user)
    assert circuits_of(imported) == {"A": 1, "B": 1, "C": None}


def test_excel_lone_circuit_number_is_dropped(other_user):
    book = Workbook()
    sheet = book.active
    sheet.title = excel.SHEET_TITLE
    sheet.append(
        ["Дата", "Вид спорта", "Длительность, мин", "Упражнение", "Вес, кг", "Повторы", "Круг"]
    )
    sheet.append(["01.09.2026", "Силовая", 60, "Одинокое", 50, 10, 2])
    sheet.append(["01.09.2026", "Силовая", 60, "Другое", 50, 10, None])
    buffer = io.BytesIO()
    book.save(buffer)

    excel_import.import_workbook(other_user, io.BytesIO(buffer.getvalue()))

    imported = Workout.objects.get(user=other_user)
    assert set(imported.sets.values_list("circuit", flat=True)) == {None}


# ---------- Бюджеты ----------


@pytest.fixture
def triset_live(user):
    workout = WorkoutFactory(user=user, duration_min=None)
    for name in "ABC":
        make(workout, ExerciseFactory(name=name), count=3, done=False, circuit=1)
    return workout


def test_live_screen_with_circuit_stays_in_budget(
    client, user, triset_live, django_assert_max_num_queries
):
    client.force_login(user)

    with django_assert_max_num_queries(9 + SIDEBAR_QUERIES):
        client.get(reverse("workout_live", args=[triset_live.pk]))


def test_summary_with_circuit_stays_in_budget(client, user, django_assert_max_num_queries):
    workout = WorkoutFactory(user=user)
    for name in "ABC":
        make(workout, ExerciseFactory(name=name), count=3, circuit=1)
    client.force_login(user)

    with django_assert_max_num_queries(9 + SIDEBAR_QUERIES):
        client.get(reverse("workout_summary", args=[workout.pk]))


@pytest.mark.parametrize("count", [2, 6])
def test_circuit_modal_does_not_scale(client, user, count, django_assert_max_num_queries):
    workout = WorkoutFactory(user=user, duration_min=None)
    for index in range(count):
        make(workout, ExerciseFactory(name=f"Упражнение {index}"), done=False)
    client.force_login(user)

    # Транзакция запроса, сессия, пользователь, тренировка, подходы и заметки.
    with django_assert_max_num_queries(7):
        client.get(reverse("live_circuits", args=[workout.pk]))


def test_foreign_circuit_modal_is_404(client, user, other_user):
    alien = WorkoutFactory(user=other_user, duration_min=None)
    client.force_login(user)

    assert client.get(reverse("live_circuits", args=[alien.pk])).status_code == 404
