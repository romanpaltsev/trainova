"""Экран «Экспорт и импорт»: доступ, формы загрузки истории и справочника, отчёты."""

import io
from datetime import date

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from workouts import excel_import, exercise_excel
from workouts.models import Exercise, Workout
from workouts.tests.factories import (
    ExerciseFactory,
    SportFactory,
    StrengthSetFactory,
    WorkoutFactory,
)
from workouts.tests.sheets import build_sheet

pytestmark = pytest.mark.django_db

PAGE = "data_transfer"


def upload(rows, name="история.xlsx"):
    return SimpleUploadedFile(name, build_sheet(rows).getvalue())


def test_page_requires_login(client):
    response = client.get(reverse(PAGE))

    assert response.status_code == 302
    assert reverse("account_login") in response["Location"]


def test_page_shows_own_workout_count(client, user, other_user):
    WorkoutFactory(user=user)
    WorkoutFactory(user=user)
    WorkoutFactory(user=other_user)
    client.force_login(user)

    response = client.get(reverse(PAGE))

    assert response.status_code == 200
    assert response.context["workouts_count"] == 2


def test_page_has_upload_form(client, user):
    client.force_login(user)

    content = client.get(reverse(PAGE)).content.decode()

    assert 'enctype="multipart/form-data"' in content
    assert 'type="file"' in content


def test_empty_diary_says_so(client, user):
    client.force_login(user)

    content = client.get(reverse(PAGE)).content.decode()

    assert "Пока выгружать нечего" in content


def test_upload_creates_workouts_and_shows_report(client, user):
    SportFactory(name="Силовая")
    client.force_login(user)

    response = client.post(
        reverse(PAGE),
        {
            "file": upload(
                [
                    {
                        "date": date(2026, 9, 4),
                        "sport": "Силовая",
                        "exercise": "Жим лёжа",
                        "weight": 80,
                        "reps": 8,
                        "duration": 60,
                    }
                ]
            )
        },
    )

    assert response.status_code == 200
    assert response.context["report"].workouts == 1
    assert Workout.objects.filter(user=user).count() == 1
    assert "Что получилось" in response.content.decode()


def test_upload_without_file_shows_form_error(client, user):
    client.force_login(user)

    response = client.post(reverse(PAGE), {})

    assert response.status_code == 200
    assert "Выберите файл." in response.content.decode()
    assert not Workout.objects.exists()


def test_upload_of_foreign_format_is_explained(client, user):
    client.force_login(user)

    response = client.post(
        reverse(PAGE), {"file": SimpleUploadedFile("заметки.txt", b"not a workbook")}
    )

    assert response.status_code == 200
    assert ".xlsx" in response.content.decode()


def test_broken_workbook_does_not_give_500(client, user):
    """Файл с правильным расширением, но не книга — человеческий текст, не 500."""
    client.force_login(user)

    response = client.post(
        reverse(PAGE), {"file": SimpleUploadedFile("история.xlsx", b"PK\x03\x04 broken")}
    )

    assert response.status_code == 200
    assert "Не получилось открыть файл" in response.content.decode()


def test_too_big_file_is_rejected(client, user, monkeypatch):
    monkeypatch.setattr(excel_import, "MAX_UPLOAD_BYTES", 10)
    client.force_login(user)

    response = client.post(reverse(PAGE), {"file": SimpleUploadedFile("история.xlsx", b"x" * 100)})

    assert "Файл больше" in response.content.decode()
    assert not Workout.objects.exists()


def test_export_link_is_on_the_page(client, user):
    workout = WorkoutFactory(user=user)
    StrengthSetFactory(workout=workout)
    client.force_login(user)

    content = client.get(reverse(PAGE)).content.decode()

    assert reverse("workout_export") in content


def test_export_of_foreign_user_is_not_reachable(client, user, other_user):
    """Выгрузка отдаёт только свои тренировки — чужих в файле нет."""
    theirs = WorkoutFactory(user=other_user)
    StrengthSetFactory(workout=theirs)
    client.force_login(user)

    response = client.get(reverse("workout_export"))
    content = b"".join(response.streaming_content)

    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content)).worksheets[0]
    assert sheet.max_row == 1


# ---------- Справочник упражнений ----------

EXERCISES = "exercise_import"


def exercises_upload(rows, name="справочник.xlsx"):
    stream = build_sheet(
        rows, sheet_title=exercise_excel.SHEET_TITLE, columns=exercise_excel.COLUMNS
    )
    return SimpleUploadedFile(name, stream.getvalue())


def test_page_offers_exercise_export_and_upload(client, user):
    client.force_login(user)

    content = client.get(reverse(PAGE)).content.decode()

    assert reverse("exercise_export") in content
    assert f'action="{reverse(EXERCISES)}#exercise-upload"' in content
    assert 'name="exercises-file"' in content


def test_two_file_inputs_have_distinct_ids(client, user):
    """Две формы на странице: у полей свои id, иначе подпись одной открывала бы другую."""
    client.force_login(user)

    content = client.get(reverse(PAGE)).content.decode()

    assert content.count('id="id_file"') == 1
    assert content.count('id="id_exercises-file"') == 1


def test_exercise_upload_changes_own_exercise_and_shows_report(client, user):
    own = ExerciseFactory(name="Жим", owner=user)
    client.force_login(user)

    response = client.post(
        reverse(EXERCISES),
        {"exercises-file": exercises_upload([{"id": own.pk, "name": "Жим узкий"}])},
    )

    own.refresh_from_db()
    assert response.status_code == 200
    assert own.name == "Жим узкий"
    assert response.context["exercise_report"].updated == 1
    assert "Жим → Жим узкий" in response.content.decode()


def test_history_form_keeps_its_address_after_exercise_upload(client, user):
    """После загрузки справочника страница живёт на его адресе, а форма истории
    обязана по-прежнему отправляться на свой."""
    client.force_login(user)

    content = client.post(
        reverse(EXERCISES), {"exercises-file": exercises_upload([{"name": "Новое"}])}
    ).content.decode()

    assert f'action="{reverse(PAGE)}"' in content


def test_exercise_upload_without_file_shows_form_error(client, user):
    client.force_login(user)

    response = client.post(reverse(EXERCISES), {})

    assert response.status_code == 200
    assert "Выберите файл." in response.content.decode()
    assert response.context["exercise_form"].errors


def test_broken_exercise_workbook_does_not_give_500(client, user):
    client.force_login(user)
    broken = SimpleUploadedFile("справочник.xlsx", b"PK\x03\x04 broken")

    response = client.post(reverse(EXERCISES), {"exercises-file": broken})

    assert response.status_code == 200
    assert "Не получилось открыть файл" in response.content.decode()


def test_history_file_in_exercise_form_is_explained(client, user):
    client.force_login(user)
    history = upload([{"date": date(2026, 9, 1), "sport": "Силовая", "exercise": "Жим"}])

    response = client.post(reverse(EXERCISES), {"exercises-file": history})

    assert "файл с историей тренировок" in response.content.decode()
    assert not Exercise.objects.exists()


def test_exercise_import_page_is_not_a_page(client, user):
    """GET по адресу загрузки — обратно на экран обмена: формы там нет."""
    client.force_login(user)

    response = client.get(reverse(EXERCISES))

    assert response.status_code == 302
    assert response["Location"] == reverse(PAGE)


def test_exercise_import_requires_login(client):
    response = client.post(reverse(EXERCISES), {})

    assert response.status_code == 302
    assert reverse("account_login") in response["Location"]
