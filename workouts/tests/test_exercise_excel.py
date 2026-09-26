"""Справочник упражнений в Excel: выгрузка и загрузка обратно.

Правило одно с экраном упражнения: своё правится целиком, у общего — только
личный шаг веса, новая строка — своё упражнение, удалить через файл нельзя.
Главный тест — круговой: нетронутый файл не меняет ничего.
"""

import io
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import unquote

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from workouts import excel, exercise_excel
from workouts.models import Exercise, ExerciseSettings
from workouts.tests.factories import (
    ExerciseFactory,
    ExerciseSettingsFactory,
    StrengthSetFactory,
    TimeSetFactory,
    WorkoutFactory,
)
from workouts.tests.sheets import build_sheet, dump_exercises

pytestmark = pytest.mark.django_db

MEASURE = Exercise.Measurement
KEYS = [column.key for column in exercise_excel.COLUMNS]


def sheet_of(user):
    """Лист справочника, собранный напрямую — без HTTP."""
    return exercise_excel.build_workbook(user)[exercise_excel.SHEET_TITLE]


def rows_of(sheet):
    """Строки листа словарями по ключам COLUMNS."""
    return [
        dict(zip(KEYS, values, strict=True))
        for values in sheet.iter_rows(min_row=2, values_only=True)
    ]


def exported(user):
    """Выгрузка пользователя как поток — ровно то, что он скачал бы."""
    buffer = io.BytesIO()
    exercise_excel.build_workbook(user).save(buffer)
    buffer.seek(0)
    return buffer


def run(user, rows, **kwargs):
    stream = build_sheet(
        rows, sheet_title=exercise_excel.SHEET_TITLE, columns=exercise_excel.COLUMNS, **kwargs
    )
    return exercise_excel.import_workbook(user, stream)


def row(exercise=None, **values):
    """Строка файла: по умолчанию — упражнение как есть, с его ID."""
    base = {}
    if exercise is not None:
        base = {
            "id": exercise.pk,
            "name": exercise.name,
            "muscle_group": exercise.muscle_group,
            "equipment": exercise.equipment,
            "measurement": exercise.get_measurement_display(),
        }
    return base | values


# ---------- Выгрузка ----------


def test_export_requires_login(client):
    response = client.get(reverse("exercise_export"))

    assert response.status_code == 302
    assert reverse("account_login") in response["Location"]


def test_export_returns_xlsx_with_dated_filename(client, user):
    client.force_login(user)

    response = client.get(reverse("exercise_export"))
    content = b"".join(response.streaming_content)

    assert response["Content-Type"] == excel.CONTENT_TYPE
    assert "attachment" in response["Content-Disposition"]
    # Имя — по-русски, как у выгрузки истории: Django кодирует его по RFC 5987.
    assert "Справочник упражнений" in unquote(response["Content-Disposition"])
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == [exercise_excel.SHEET_TITLE, excel.HELP_SHEET_TITLE]


def test_export_filename_is_dated():
    name = exercise_excel.export_filename(date(2026, 9, 26))

    assert name == "Справочник упражнений 2026-09-26.xlsx"


def test_header_matches_columns(user):
    titles = [cell.value for cell in sheet_of(user)[1]]

    assert titles == [column.title for column in exercise_excel.COLUMNS]


def test_export_has_global_and_own_but_not_foreign(user, other_user):
    """Изоляция: чужие личные упражнения в файл не попадают."""
    ExerciseFactory(name="Жим штанги лёжа")
    ExerciseFactory(name="Мой жим", owner=user)
    ExerciseFactory(name="Чужой жим", owner=other_user)

    names = {line["name"] for line in rows_of(sheet_of(user))}

    assert names == {"Жим штанги лёжа", "Мой жим"}


def test_owner_column_says_whose(user):
    ExerciseFactory(name="Общее")
    ExerciseFactory(name="Своё", owner=user)

    owners = {line["name"]: line["owner"] for line in rows_of(sheet_of(user))}

    assert owners == {"Общее": "общее", "Своё": "моё"}


def test_weight_step_column(user):
    """Свой шаг, умолчание 2,5 и пусто у упражнения без веса — текстом, как на экране."""
    dumbbell = ExerciseFactory(name="Гантели")
    ExerciseSettingsFactory(user=user, exercise=dumbbell, weight_step=Decimal("1.25"))
    ExerciseFactory(name="Штанга")
    ExerciseFactory(name="Планка", measurement=MEASURE.TIME)

    steps = {line["name"]: line["weight_step"] for line in rows_of(sheet_of(user))}

    assert steps == {"Гантели": "1,25", "Штанга": "2,5", "Планка": None}


def test_steps_of_another_user_do_not_leak(user, other_user):
    exercise = ExerciseFactory(name="Гантели")
    ExerciseSettingsFactory(user=other_user, exercise=exercise, weight_step=Decimal("0.5"))

    (line,) = rows_of(sheet_of(user))

    assert line["weight_step"] == "2,5"


def test_workouts_column_counts_only_own_finished_workouts(user, other_user):
    bench = ExerciseFactory(name="Жим")
    finished = WorkoutFactory(user=user)
    StrengthSetFactory(workout=finished, exercise=bench, set_number=1)
    StrengthSetFactory(workout=finished, exercise=bench, set_number=2)
    draft = WorkoutFactory(user=user, started_at=None, duration_min=None)
    StrengthSetFactory(workout=draft, exercise=bench, set_number=1, done=False)
    StrengthSetFactory(workout=WorkoutFactory(user=other_user), exercise=bench, set_number=1)

    (line,) = rows_of(sheet_of(user))

    assert line["workouts"] == 1


def test_rows_are_ordered_by_group_then_name(user):
    ExerciseFactory(name="Тяга", muscle_group="Спина")
    ExerciseFactory(name="Разведение", muscle_group="Грудь")
    ExerciseFactory(name="Жим", muscle_group="Грудь")
    ExerciseFactory(name="Прыжки", muscle_group="")

    names = [line["name"] for line in rows_of(sheet_of(user))]

    assert names == ["Жим", "Разведение", "Тяга", "Прыжки"]


def test_formula_like_name_stays_text(user):
    ExerciseFactory(name="=2+2", owner=user)

    cell = sheet_of(user)["B2"]

    assert cell.value == "=2+2"
    assert cell.data_type == "s"


def test_measurement_column_suggests_units(user):
    ExerciseFactory()

    (validation,) = sheet_of(user).data_validations.dataValidation

    assert validation.type == "list"
    assert "Вес × повторы" in validation.formula1
    # Подсказка, а не запрет: набранное руками разбор поймёт сам.
    assert not validation.showErrorMessage


# ---------- Круговой ----------


def test_round_trip_changes_nothing(user):
    """Главный тест: выгрузил и загрузил нетронутым — ни одной записи в базу.

    Набор нарочно неудобный: своё упражнение с именем общего (так их развела
    миграция точных имён), своё «грудь» при общем «Грудь», удержание без веса,
    строка настроек с умолчанием и имя с двойным пробелом.
    """
    ExerciseFactory(name="Жим штанги лёжа", muscle_group="Грудь")
    ExerciseFactory(name="Жим штанги лёжа", owner=user, muscle_group="грудь")
    ExerciseFactory(name="Планка", measurement=MEASURE.TIME, muscle_group="Пресс")
    own = ExerciseFactory(name="Тяга  к поясу", owner=user, muscle_group="Спина")
    ExerciseSettingsFactory(user=user, exercise=own, weight_step=Decimal("2.5"))
    before = dump_exercises(user)

    with CaptureQueriesContext(connection) as queries:
        report = exercise_excel.import_workbook(user, exported(user))

    sql = [query["sql"] for query in queries.captured_queries]
    assert not [text for text in sql if text.startswith(("UPDATE", "INSERT"))]
    assert report.unchanged == 4
    assert (report.created, report.updated, report.steps) == ([], 0, 0)
    assert (report.errors, report.warnings) == ([], [])
    assert dump_exercises(user) == before


def test_saved_file_undoes_edits(user):
    """Сохранённая до правок выгрузка возвращает справочник как было."""
    own = ExerciseFactory(name="Жим", owner=user, muscle_group="Грудь")
    backup = exported(user)
    run(user, [row(own, name="Жим узкий", muscle_group="Руки")])

    exercise_excel.import_workbook(user, backup)

    own.refresh_from_db()
    assert (own.name, own.muscle_group) == ("Жим", "Грудь")


# ---------- Свои упражнения ----------


def test_rename_by_id_keeps_history(user):
    own = ExerciseFactory(name="Жим в тренажёре", owner=user)
    done = StrengthSetFactory(workout=WorkoutFactory(user=user), exercise=own)

    report = run(user, [row(own, name="Жим от груди в тренажёре")])

    own.refresh_from_db()
    done.refresh_from_db()
    assert own.name == "Жим от груди в тренажёре"
    assert done.exercise_id == own.pk
    assert report.renamed == [("Жим в тренажёре", "Жим от груди в тренажёре")]
    assert report.updated == 1


@pytest.mark.parametrize("taken_by", ["own", "global"])
def test_rename_to_a_taken_name_is_rejected(user, taken_by):
    own = ExerciseFactory(name="Жим", owner=user)
    ExerciseFactory(name="Тяга", owner=user if taken_by == "own" else None)

    report = run(user, [row(own, name="тяга", muscle_group="Спина")])

    own.refresh_from_db()
    assert (own.name, own.muscle_group) == ("Жим", "Грудь")
    assert "уже занято" in report.errors[0]


def test_case_only_rename_is_allowed(user):
    own = ExerciseFactory(name="жим лёжа", owner=user)

    run(user, [row(own, name="Жим лёжа")])

    own.refresh_from_db()
    assert own.name == "Жим лёжа"


def test_row_without_id_finds_by_name_and_keeps_its_case(user):
    """Без ID название — ключ поиска: другой регистр значит «найди», а не «переименуй»."""
    own = ExerciseFactory(name="Жим лёжа", owner=user, muscle_group="Грудь")

    report = run(user, [{"name": "жим лёжа", "muscle_group": "Руки"}])

    own.refresh_from_db()
    assert (own.name, own.muscle_group) == ("Жим лёжа", "Руки")
    assert report.created == []
    assert Exercise.objects.filter(owner=user).count() == 1


def test_renaming_own_twin_of_global_keeps_global_name_taken(user):
    """Своё «Жим» рядом с общим «Жим»: переименование своего не освобождает имя общего."""
    ExerciseFactory(name="Жим")
    twin = ExerciseFactory(name="Жим", owner=user)
    other = ExerciseFactory(name="Тяга", owner=user)

    report = run(user, [row(twin, name="Жим узким хватом"), row(other, name="Жим")])

    twin.refresh_from_db()
    other.refresh_from_db()
    assert twin.name == "Жим узким хватом"
    assert other.name == "Тяга"
    assert len(report.errors) == 1


def test_group_and_equipment_follow_existing_spelling(user):
    ExerciseFactory(name="Общее", muscle_group="Грудь", equipment="Гантели")
    own = ExerciseFactory(name="Своё", owner=user, muscle_group="", equipment="")

    run(user, [row(own, muscle_group="грудь", equipment="  гантели ")])

    own.refresh_from_db()
    assert (own.muscle_group, own.equipment) == ("Грудь", "Гантели")


def test_new_spelling_in_a_file_is_used_by_later_rows(user):
    first = ExerciseFactory(name="Первое", owner=user, muscle_group="")
    second = ExerciseFactory(name="Второе", owner=user, muscle_group="")

    run(user, [row(first, muscle_group="Кор"), row(second, muscle_group="кор")])

    second.refresh_from_db()
    assert second.muscle_group == "Кор"


def test_empty_facet_cell_clears_the_value(user):
    own = ExerciseFactory(name="Жим", owner=user, muscle_group="Грудь", equipment="Штанга")

    run(user, [row(own, muscle_group="", equipment="")])

    own.refresh_from_db()
    assert (own.muscle_group, own.equipment) == ("", "")


def test_missing_facet_column_keeps_the_value(user):
    own = ExerciseFactory(name="Жим", owner=user, muscle_group="Грудь", equipment="Штанга")

    report = run(user, [{"id": own.pk, "name": "Жим"}], header=["ID", "Упражнение"])

    own.refresh_from_db()
    assert (own.muscle_group, own.equipment) == ("Грудь", "Штанга")
    assert report.unchanged == 1


@pytest.mark.parametrize(
    ("text", "measurement"),
    [
        pytest.param("Повторы", MEASURE.REPS, id="как в файле"),
        pytest.param("вес x повторы", MEASURE.WEIGHT_REPS, id="латинская x"),
        pytest.param("ВЕС Х ПОВТОРЫ", MEASURE.WEIGHT_REPS, id="кириллическая х"),
        pytest.param("время+вес", MEASURE.TIME_WEIGHT, id="плюс без пробелов"),
    ],
)
def test_measurement_labels_are_understood(user, text, measurement):
    own = ExerciseFactory(name="Своё", owner=user, measurement=MEASURE.TIME)

    run(user, [row(own, measurement=text)])

    own.refresh_from_db()
    assert own.measurement == measurement


def test_unknown_measurement_rejects_the_row(user):
    own = ExerciseFactory(name="Своё", owner=user)

    report = run(user, [row(own, name="Переименованное", measurement="килограммы")])

    own.refresh_from_db()
    assert own.name == "Своё"
    assert "Измерение — одно из" in report.errors[0]


def test_measurement_change_is_explained(user):
    own = ExerciseFactory(name="Вис", owner=user)

    report = run(user, [row(own, measurement="Время")])

    own.refresh_from_db()
    assert own.measurement == MEASURE.TIME
    assert any("остались в прежних единицах" in text for text in report.warnings)


def test_new_row_creates_own_exercise(user):
    ExerciseFactory(name="Общее", muscle_group="Спина")

    report = run(
        user,
        [{"name": "Тяга резины", "muscle_group": "спина", "measurement": "Повторы"}],
    )

    created = Exercise.objects.get(name="Тяга резины")
    assert created.owner == user
    assert (created.muscle_group, created.measurement) == ("Спина", MEASURE.REPS)
    assert report.created == ["Тяга резины"]


def test_new_row_without_unit_counts_weight_and_reps(user):
    run(user, [{"name": "Новое"}], header=["Упражнение"])

    assert Exercise.objects.get(name="Новое").measurement == MEASURE.WEIGHT_REPS


def test_same_new_name_twice_creates_one_exercise(user):
    report = run(user, [{"name": "Новое"}, {"name": "новое"}], header=["Упражнение"])

    assert Exercise.objects.filter(owner=user).count() == 1
    assert "уже было в строке 2" in report.errors[0]


# ---------- Общие упражнения ----------


def test_global_row_applies_only_the_weight_step(user):
    shared = ExerciseFactory(name="Жим штанги лёжа", muscle_group="Грудь")

    report = run(user, [row(shared, name="Жим лёжа", muscle_group="Руки", weight_step="1,25")])

    shared.refresh_from_db()
    assert (shared.name, shared.muscle_group) == ("Жим штанги лёжа", "Грудь")
    assert ExerciseSettings.objects.get(user=user, exercise=shared).weight_step == Decimal("1.25")
    assert report.steps == 1
    assert "Жим штанги лёжа" in report.warnings[0]


def test_global_row_found_by_name_in_another_case_is_not_an_edit(user):
    ExerciseFactory(name="Жим штанги лёжа", muscle_group="Грудь")

    report = run(user, [{"name": "жим штанги лёжа", "muscle_group": "Грудь"}])

    assert report.warnings == []
    assert report.unchanged == 1


# ---------- Шаг веса ----------


def test_default_step_does_not_create_settings(user):
    own = ExerciseFactory(name="Жим", owner=user)

    report = run(user, [row(own, weight_step="2,5")])

    assert not ExerciseSettings.objects.exists()
    assert report.unchanged == 1


@pytest.mark.parametrize("value", [1.25, "1,25", "1.25"], ids=["число", "запятая", "точка"])
def test_step_is_read_as_number_and_as_text(user, value):
    own = ExerciseFactory(name="Жим", owner=user)

    run(user, [row(own, weight_step=value)])

    assert ExerciseSettings.objects.get(user=user).weight_step == Decimal("1.25")


def test_changed_step_of_exercise_without_weight_is_a_warning(user):
    plank = ExerciseFactory(name="Планка", owner=user, measurement=MEASURE.TIME)

    report = run(user, [row(plank, weight_step="1,25")])

    assert not ExerciseSettings.objects.exists()
    assert report.errors == []
    assert "нет веса" in report.warnings[0]


def test_old_step_left_after_unit_change_is_not_noise(user):
    """Сменили единицу на время, а «2,5» в строке осталось — это не правка шага."""
    own = ExerciseFactory(name="Вис", owner=user)

    report = run(user, [row(own, measurement="Время", weight_step="2,5")])

    assert not any("нет веса" in text for text in report.warnings)


@pytest.mark.parametrize(
    "value",
    ["0,1", "nan", "много", datetime(2026, 5, 1)],
    ids=["меньше мин", "nan", "мусор", "дата"],
)
def test_bad_step_rejects_the_whole_row(user, value):
    """Строка применяется целиком или никак: плохой шаг не пропустит переименование."""
    own = ExerciseFactory(name="Жим", owner=user)

    report = run(user, [row(own, name="Жим узкий", weight_step=value)])

    own.refresh_from_db()
    assert own.name == "Жим"
    assert not ExerciseSettings.objects.exists()
    assert report.error_count == 1


def test_step_turned_into_date_gets_a_hint(user):
    own = ExerciseFactory(name="Жим", owner=user)

    report = run(user, [row(own, weight_step=datetime(2026, 5, 1))])

    assert "превратил шаг веса в дату" in report.errors[0]


# ---------- ID и ошибки строк ----------


def test_foreign_id_is_rejected_and_left_untouched(user, other_user):
    """Изоляция: чужое упражнение по ID не правится и не копируется."""
    foreign = ExerciseFactory(name="Чужое", owner=other_user)
    before = dump_exercises(other_user)

    report = run(user, [row(foreign, name="Моё теперь", weight_step="1")])

    assert dump_exercises(other_user) == before
    assert not Exercise.objects.filter(owner=user).exists()
    assert not ExerciseSettings.objects.filter(user=user).exists()
    assert f"ID {foreign.pk} нет" in report.errors[0]


def test_deleted_id_is_not_resurrected(user):
    report = run(user, [{"id": 999_999, "name": "Удалённое"}])

    assert not Exercise.objects.filter(name="Удалённое").exists()
    assert "очистите ID" in report.errors[0]


@pytest.mark.parametrize("value", [True, 5.7, "abc", 0], ids=["ИСТИНА", "дробь", "буквы", "ноль"])
def test_bad_id_rejects_the_row(user, value):
    report = run(user, [{"id": value, "name": "Упражнение"}])

    assert "ID — это номер из выгрузки" in report.errors[0]
    assert not Exercise.objects.exists()


def test_integer_float_id_is_accepted(user):
    own = ExerciseFactory(name="Жим", owner=user)

    run(user, [row(own, id=float(own.pk), name="Жим узкий")])

    own.refresh_from_db()
    assert own.name == "Жим узкий"


def test_same_exercise_twice_is_reported(user):
    own = ExerciseFactory(name="Жим", owner=user)

    report = run(user, [row(own, muscle_group="Руки"), row(own, muscle_group="Спина")])

    own.refresh_from_db()
    assert own.muscle_group == "Руки"
    assert "уже было в строке 2" in report.errors[0]


def test_row_without_name_is_rejected(user):
    own = ExerciseFactory(name="Жим", owner=user)

    report = run(user, [row(own, name="")])

    own.refresh_from_db()
    assert own.name == "Жим"
    assert "Нет названия" in report.errors[0]


def test_too_long_name_is_an_error_not_a_crash(user):
    report = run(user, [{"name": "Ж" * 81}], header=["Упражнение"])

    assert "длиннее 80" in report.errors[0]
    assert not Exercise.objects.exists()


def test_error_rows_do_not_block_the_rest(user):
    good = ExerciseFactory(name="Жим", owner=user)

    report = run(user, [{"id": "abc", "name": "Плохое"}, row(good, muscle_group="Руки")])

    good.refresh_from_db()
    assert good.muscle_group == "Руки"
    assert report.error_count == 1
    assert report.errors[0].startswith("Строка 2:")


def test_import_does_not_touch_other_users_data(user, other_user):
    shared = ExerciseFactory(name="Общее")
    ExerciseFactory(name="Чужое", owner=other_user)
    ExerciseSettingsFactory(user=other_user, exercise=shared, weight_step=Decimal("5"))
    before = dump_exercises(other_user)

    run(user, [row(shared, weight_step="1"), {"name": "Чужое", "muscle_group": "Спина"}])

    assert dump_exercises(other_user) == before
    # «Чужое» у пользователя — новое своё упражнение, а не чужое найденное.
    assert Exercise.objects.get(owner=user, name="Чужое").muscle_group == "Спина"


# ---------- Файл целиком ----------


def test_missing_name_column_gives_human_error(user):
    with pytest.raises(excel.WorkbookError, match="не хватает колонок: Упражнение"):
        run(user, [{"id": 1}], header=["ID", "Группа мышц"])


def test_history_file_is_recognised(user):
    history = build_sheet([{"date": date(2026, 9, 1), "sport": "Силовая", "exercise": "Жим"}])

    with pytest.raises(excel.WorkbookError, match="файл с историей тренировок"):
        exercise_excel.import_workbook(user, history)
    assert not Exercise.objects.exists()


def test_too_many_rows_are_rejected(user, monkeypatch):
    monkeypatch.setattr(exercise_excel, "MAX_ROWS", 2)

    with pytest.raises(excel.WorkbookError, match="Слишком много строк"):
        run(user, [{"name": f"Новое {n}"} for n in range(3)], header=["Упражнение"])


def test_not_a_workbook_gives_human_error(user):
    with pytest.raises(excel.WorkbookError, match="Не получилось открыть файл"):
        exercise_excel.import_workbook(user, io.BytesIO(b"PK\x03\x04 broken"))


def test_time_set_history_is_not_affected_by_unit_change(user):
    """Смена единицы не переписывает записанные подходы: у них свой снимок."""
    plank = ExerciseFactory(name="Планка", owner=user, measurement=MEASURE.TIME)
    done = TimeSetFactory(workout=WorkoutFactory(user=user), exercise=plank)

    run(user, [row(plank, measurement="Время + вес")])

    done.refresh_from_db()
    assert done.measurement == MEASURE.TIME
