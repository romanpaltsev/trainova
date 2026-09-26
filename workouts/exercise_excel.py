"""Справочник упражнений в Excel: выгрузка и загрузка обратно.

Файл отдельный от истории: строка здесь — упражнение, а не подход, и колонки
истории — договор, который ломать нельзя. Колонки справочника объявлены здесь
один раз (COLUMNS), и разбор берёт их отсюда же — иначе скачанный файл перестал
бы загружаться обратно.

Правило то же, что на странице упражнения: своё упражнение правится целиком —
название, группа мышц, снаряд, единица; общее правит только администратор, и у
него файл меняет лишь личный шаг веса (ExerciseSettings). Новая строка — новое
своё упражнение. Удалить через файл нельзя ничего: строка, пропавшая из таблицы,
чаще случайность, чем решение, а упражнение из записанной тренировки не удалить
и на экране.

Повторная загрузка неизменённого файла не пишет ничего: каждое значение
сравнивается с текущим, и запись идёт, только если они различаются. На этом же
держится отмена: сохранённый заранее файл возвращает справочник как было.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal

from django.db import DataError, IntegrityError, transaction
from django.db.models import Count
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from workouts import excel
from workouts.models import (
    DEFAULT_WEIGHT_STEP,
    EQUIPMENT_MAX_LENGTH,
    EXERCISE_NAME_MAX_LENGTH,
    MEASUREMENT_FIELDS,
    MUSCLE_GROUP_MAX_LENGTH,
    Exercise,
    ExerciseSettings,
    collapse_spaces,
    decimal_display,
    exercise_usage,
    facets_from_pairs,
    normalize_facet,
    parse_weight_step,
)

SHEET_TITLE = "Упражнения"
OWN_LABEL = "моё"
GLOBAL_LABEL = "общее"

# Единственный источник правды о формате справочника. «Чьё» и «Тренировок» —
# только для человека: разбор их не читает, поэтому поменять в них что-то нельзя.
COLUMNS: tuple[excel.Column, ...] = (
    excel.Column("id", "ID", 8, "0"),
    excel.Column("name", "Упражнение", 40),
    excel.Column("muscle_group", "Группа мышц", 16),
    excel.Column("equipment", "Снаряд", 14),
    excel.Column("measurement", "Измерение", 16),
    # Текстом «2,5», как на экране, а не числом: русский Excel превращает
    # набранное «1.5» в дату 1 мая, а текстовая ячейка оставляет его как есть.
    excel.Column("weight_step", "Шаг веса, кг", 13, "@"),
    excel.Column("owner", "Чьё", 10),
    excel.Column("workouts", "Тренировок", 12, "0"),
)
# Без ID таблица законна: так её собирают руками, и строки ищутся по названию.
REQUIRED_TITLES = ("Упражнение",)
# Справочник — это десятки строк, а не тысячи, как история. Предел держит
# загрузку далеко от таймаута запроса: изменённая строка — до четырёх запросов.
MAX_ROWS = 1_000
# Сколько названий перечислять в одном предупреждении — дальше «и ещё N».
NAMES_SHOWN = 10
# Сколько пустых строк под таблицей получают подсказку-список в «Измерении»:
# новые упражнения дописывают снизу.
SPARE_ROWS = 200

MEASUREMENT_LABELS = [label for _, label in Exercise.Measurement.choices]

HELP_LINES = (
    "Как заполнять таблицу",
    "",
    "Одна строка — одно упражнение: и общие упражнения приложения, и ваши собственные.",
    "",
    "«ID» не меняйте: по нему приложение узнаёт упражнение, даже если его переименовали.",
    "Для нового упражнения оставьте «ID» пустым — оно появится в справочнике как ваше.",
    "Строка без «ID» ищется по названию, регистр букв не важен.",
    "",
    "Свои упражнения («моё» в колонке «Чьё») меняются целиком: название, группа мышц,",
    "снаряд и измерение. Общие правит администратор — у них загрузка меняет только ваш",
    "шаг веса.",
    "",
    "«Измерение» — одно из: " + ", ".join(MEASUREMENT_LABELS) + ".",
    "«Шаг веса, кг» — на сколько кнопки «−» и «+» меняют вес, от 0,25 до 50 кг; дробь",
    "пишите через запятую: 1,25. Упражнениям без веса шаг не нужен.",
    "Пустая ячейка измерения или шага ничего не меняет.",
    "",
    "Пустая ячейка группы мышц или снаряда убирает значение. Колонку, которой нет в файле,",
    "загрузка не трогает — свою таблицу можно собрать хоть из одной колонки «Упражнение».",
    "",
    "Удалённая из файла строка ничего не удаляет: убрать своё упражнение можно на его",
    "странице в приложении.",
    "",
    "«Чьё» и «Тренировок» — для справки, при загрузке они не читаются.",
    "",
    "Файл можно загружать повторно: то, что уже совпадает, останется как есть. Поэтому",
    "сохраните скачанный файл до правок: его загрузка вернёт справочник как было",
    "(добавленные упражнения останутся).",
    "Таблицу из другого аккаунта загружайте с пустой колонкой «ID».",
)


# --- Выгрузка ------------------------------------------------------------------


def sheet_rows(user):
    """Строки листа: весь видимый справочник — два запроса, сколько бы его ни было."""
    exercises = Exercise.objects.visible_to(user).annotate(
        workouts_count=Count("sets__workout", distinct=True, filter=exercise_usage(user))
    )
    steps = dict(
        ExerciseSettings.objects.filter(user=user).values_list("exercise_id", "weight_step")
    )
    rows = []
    for exercise in sorted(exercises, key=_order):
        has_weight = "weight_kg" in MEASUREMENT_FIELDS[exercise.measurement]
        step = steps.get(exercise.pk, DEFAULT_WEIGHT_STEP)
        rows.append(
            {
                "id": exercise.pk,
                "name": exercise.name,
                "muscle_group": exercise.muscle_group,
                "equipment": exercise.equipment,
                "measurement": exercise.get_measurement_display(),
                "weight_step": decimal_display(step) if has_weight else "",
                "owner": GLOBAL_LABEL if exercise.is_global else OWN_LABEL,
                "workouts": exercise.workouts_count,
            }
        )
    return rows


def _order(exercise):
    """Как в каталоге: по группе мышц, упражнения без группы — в конце."""
    return (not exercise.muscle_group, exercise.muscle_group.casefold(), exercise.name.casefold())


def build_workbook(user):
    """Книга целиком: лист со справочником и лист с подсказкой."""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE
    rows = sheet_rows(user)
    excel.write_table(sheet, COLUMNS, rows)
    _measurement_dropdown(sheet, len(rows))
    excel.write_help(book, HELP_LINES)
    return book


def _measurement_dropdown(sheet, count):
    """Выпадающий список единиц в «Измерении».

    Только подсказка, не запрет (showErrorMessage=False): набранное руками «вес x
    повторы» разбор всё равно поймёт, и спорить с человеком окном Excel незачем.
    """
    position = next(i for i, column in enumerate(COLUMNS, start=1) if column.key == "measurement")
    letter = get_column_letter(position)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(MEASUREMENT_LABELS) + '"',
        allow_blank=True,
        showErrorMessage=False,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{count + 1 + SPARE_ROWS}")


def export_filename(today):
    """«Справочник упражнений 2026-09-26.xlsx» — рядом с выгрузкой истории не спутать."""
    return f"Справочник упражнений {today:%Y-%m-%d}.xlsx"


# --- Разбор строк --------------------------------------------------------------


@dataclass
class ExerciseRow:
    """Строка листа, приведённая к типам модели.

    None у группы мышц и снаряда значит «колонки нет в файле — не трогать», а
    пустая строка — «колонка есть, ячейка пустая — убрать значение». У единицы и
    шага пусто и нет — одно и то же: не менять.
    """

    number: int
    exercise_id: int | None = None
    name: str = ""
    muscle_group: str | None = None
    equipment: str | None = None
    measurement: str | None = None
    weight_step: Decimal | None = None
    errors: list[str] = field(default_factory=list)


def _measurement_key(text):
    """Подпись единицы без мелких различий: «вес x повторы» — это «Вес × повторы»."""
    text = text.lower()
    # Латинская x, звёздочка и знак умножения — где угодно; кириллическая «х» —
    # только отдельным словом, иначе досталось бы буквам внутри слов.
    text = re.sub(r"\s*[×x*]\s*", " × ", text)
    text = re.sub(r"(?<=\s)х(?=\s)", "×", text)
    text = re.sub(r"\s*\+\s*", " + ", text)
    return " ".join(text.split())


MEASUREMENTS = {_measurement_key(label): value for value, label in Exercise.Measurement.choices}


def _cell_id(value):
    """ID из выгрузки: целое число. TRUE, 5,7 или «abc» — не ID, а ошибка строки."""
    if value is None or value == "":
        return None
    # bool — подкласс int, и ИСТИНА из Excel стала бы упражнением с ID 1.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = excel.cell_text(value) if not isinstance(value, bool) else "?"
    if not text:
        return None
    if not text.isdigit() or int(text) == 0:
        raise ValueError(
            "ID — это номер из выгрузки. Для нового упражнения оставьте ячейку пустой."
        )
    return int(text)


def _cell_limited(value, *, limit, what):
    """Текст ячейки не длиннее поля: обрезать молча нельзя — человек не узнает."""
    text = excel.cell_text(value)
    if len(text) > limit:
        raise ValueError(f"{what} длиннее {limit} символов.")
    return text


def _cell_measurement(value):
    text = excel.cell_text(value)
    if not text:
        return None
    measurement = MEASUREMENTS.get(_measurement_key(text))
    if measurement is None:
        labels = ", ".join(f"«{label}»" for label in MEASUREMENT_LABELS)
        raise ValueError(f"Измерение — одно из: {labels}.")
    return measurement


def _cell_step(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date, time)):
        raise ValueError(
            "Excel превратил шаг веса в дату — напишите его через запятую, например 1,25."
        )
    # Число из Excel приходит float'ом: str(2.5) → «2.5», дальше общий разбор с
    # его границами и округлением до сотых. Текст «1,25» разбирается им же.
    return parse_weight_step(excel.cell_text(value))


def _exercise_row(number, values, index):
    row = ExerciseRow(number=number)
    cell = excel.cell_getter(values, index)

    def read(key, parser, target):
        try:
            setattr(row, target, parser(cell(key)))
        except ValueError as error:
            row.errors.append(str(error))

    read("id", _cell_id, "exercise_id")
    read(
        "name",
        lambda value: _cell_limited(value, limit=EXERCISE_NAME_MAX_LENGTH, what="Название"),
        "name",
    )
    # Отсутствующая колонка и пустая ячейка значат разное, поэтому смотрим в
    # индекс шапки: колонки нет — поле остаётся None, и значение не тронем.
    if index.get("muscle_group") is not None:
        read(
            "muscle_group",
            lambda value: _cell_limited(value, limit=MUSCLE_GROUP_MAX_LENGTH, what="Группа мышц"),
            "muscle_group",
        )
    if index.get("equipment") is not None:
        read(
            "equipment",
            lambda value: _cell_limited(value, limit=EQUIPMENT_MAX_LENGTH, what="Снаряд"),
            "equipment",
        )
    read("measurement", _cell_measurement, "measurement")
    read("weight_step", _cell_step, "weight_step")
    return row


def _reject_history(titles):
    """Файл с историей несёт колонку «Упражнение» и прошёл бы проверку шапки —
    а потом наплодил бы упражнений из строк-подходов. Узнаём его по датам."""
    names = {excel.cell_text(title).lower() for title in titles}
    if {"дата", "вид спорта"} <= names:
        raise excel.WorkbookError(
            "Это файл с историей тренировок, а не справочник. Его загружают выше, "
            "в блоке «История тренировок»."
        )


def read_sheet(stream):
    """Строки книги со справочником. Кидает excel.WorkbookError, если читать нечего."""
    return excel.read_rows(
        stream,
        sheet_title=SHEET_TITLE,
        columns=COLUMNS,
        required_titles=REQUIRED_TITLES,
        parse_row=_exercise_row,
        max_rows=MAX_ROWS,
        check_titles=_reject_history,
    )


# --- Загрузка ------------------------------------------------------------------


@dataclass
class ExerciseReport(excel.ReportLog):
    """Что получилось. Ошибки и предупреждения — в общем журнале excel.ReportLog."""

    created: list[str] = field(default_factory=list)
    # Пары «было → стало»: если в Excel отсортировали одну колонку, а не всю
    # таблицу, названия съедут относительно ID — список покажет это сразу.
    renamed: list[tuple[str, str]] = field(default_factory=list)
    updated: int = 0
    steps: int = 0
    unchanged: int = 0
    shared_skipped: list[str] = field(default_factory=list)
    measurement_changed: list[str] = field(default_factory=list)

    # Списки в отчёте короткие: таблицу друга на сотню строк перечислять целиком
    # незачем, число показано в итоге.
    @property
    def created_shown(self):
        return self.created[:NAMES_SHOWN]

    @property
    def created_hidden(self):
        return max(0, len(self.created) - NAMES_SHOWN)

    @property
    def renamed_shown(self):
        return self.renamed[:NAMES_SHOWN]

    @property
    def renamed_hidden(self):
        return max(0, len(self.renamed) - NAMES_SHOWN)


class RowError(Exception):
    """Строку применить нельзя; args[0] — текст для человека."""


class Catalog:
    """Видимый справочник в памяти: одно чтение на всю загрузку.

    Меняется только после того, как запись в базу удалась: откат savepoint'а
    иначе оставил бы здесь упражнение или имя, которых в базе нет, и следующие
    строки сверялись бы с выдумкой.

    По названию хранится список, а не одно упражнение: своё может называться
    так же, как общее (так их развела миграция 0034), и переименование своего не
    должно стирать из памяти общее с тем же именем.
    """

    def __init__(self, user):
        exercises = list(Exercise.objects.visible_to(user))
        self.by_pk = {exercise.pk: exercise for exercise in exercises}
        self.by_name = {}
        for exercise in exercises:
            self.by_name.setdefault(exercise.name.lower(), []).append(exercise)
        self.steps = dict(
            ExerciseSettings.objects.filter(user=user).values_list("exercise_id", "weight_step")
        )
        facets = facets_from_pairs((item.muscle_group, item.equipment) for item in exercises)
        self.groups = list(facets.muscle_groups)
        self.equipment = list(facets.equipment)

    def find(self, name):
        """Упражнение по названию без учёта регистра; своё раньше общего —
        то же правило, что у services.exercise_for_name."""
        matches = self.by_name.get(name.lower(), [])
        return min(matches, key=lambda exercise: exercise.is_global, default=None)

    def taken(self, name, *, exclude):
        """Занято ли название другим упражнением — как у ExerciseRenameView:
        по всему видимому справочнику, иначе в списке появились бы две
        одинаковые строки."""
        return any(exercise.pk != exclude for exercise in self.by_name.get(name.lower(), []))

    def step(self, exercise):
        return self.steps.get(exercise.pk, DEFAULT_WEIGHT_STEP)

    def add(self, exercise):
        self.by_pk[exercise.pk] = exercise
        self.by_name.setdefault(exercise.name.lower(), []).append(exercise)
        self.remember_facets(exercise)

    def rename(self, exercise, name):
        self.by_name[exercise.name.lower()].remove(exercise)
        exercise.name = name
        self.by_name.setdefault(name.lower(), []).append(exercise)

    def remember_facets(self, exercise):
        """Новое написание становится принятым: «кор» в следующей строке — это «Кор»."""
        if exercise.muscle_group and exercise.muscle_group not in self.groups:
            self.groups.append(exercise.muscle_group)
        if exercise.equipment and exercise.equipment not in self.equipment:
            self.equipment.append(exercise.equipment)


def import_workbook(user, stream):
    """Прочитать книгу и применить её к справочнику. Кидает excel.WorkbookError."""
    rows = read_sheet(stream)
    report = ExerciseReport()
    catalog = Catalog(user)
    seen = {}
    for row in rows:
        if row.errors:
            report.add_error(row.number, " ".join(row.errors) + " Строка пропущена.")
            continue
        try:
            _apply_row(user, catalog, report, row, seen)
        except RowError as error:
            report.add_error(row.number, str(error))
    _summarize(report)
    return report


def _apply_row(user, catalog, report, row, seen):
    """Строка применяется целиком или никак: всё, что может её отклонить,
    проверяется до записи — и ошибка в шаге не пропустит переименование."""
    exercise = _target(catalog, row)
    if exercise is not None and exercise.pk in seen:
        raise RowError(f"Это упражнение уже было в строке {seen[exercise.pk]} — повтор пропущен.")

    fields, changes = {}, {}
    if exercise is None:
        fields = _new_fields(catalog, row)
        measurement = fields["measurement"]
    elif exercise.is_global:
        if _differs_from_shared(catalog, exercise, row):
            report.shared_skipped.append(exercise.name)
        measurement = exercise.measurement
    else:
        changes = _own_changes(catalog, exercise, row)
        measurement = changes.get("measurement", exercise.measurement)
    step = _step_change(catalog, report, exercise, row, measurement)

    if exercise is not None and not changes and step is None:
        seen[exercise.pk] = row.number
        report.unchanged += 1
        return

    # Строка пишется своим savepoint'ом и только когда есть что писать: занятое
    # в эту же секунду из другой вкладки имя упрётся в ограничение базы, и
    # откатится одна эта строка, а неизменённые строки не стоят ни запроса.
    try:
        with transaction.atomic():
            if exercise is None:
                saved = Exercise.objects.create(owner=user, **fields)
            else:
                saved = exercise
                # owner в самом запросе: правило «своё — да, общее и чужое — нет»
                # держит база, а не только проверка выше.
                if changes and not Exercise.objects.filter(pk=exercise.pk, owner=user).update(
                    **changes
                ):
                    raise RowError("Упражнение пропало из справочника — строка пропущена.")
            if step is not None:
                ExerciseSettings.objects.update_or_create(
                    user=user, exercise_id=saved.pk, defaults={"weight_step": step}
                )
    except (IntegrityError, DataError):
        raise RowError("Не удалось сохранить упражнение — строка пропущена.") from None

    if exercise is None:
        catalog.add(saved)
        report.created.append(saved.name)
    elif changes:
        if "name" in changes:
            report.renamed.append((exercise.name, changes["name"]))
            catalog.rename(exercise, changes["name"])
        for name, value in changes.items():
            setattr(exercise, name, value)
        catalog.remember_facets(exercise)
        report.updated += 1
        if "measurement" in changes:
            report.measurement_changed.append(exercise.name)
    if step is not None:
        catalog.steps[saved.pk] = step
        report.steps += 1
    seen[saved.pk] = row.number


def _target(catalog, row):
    """Какое упражнение правит строка; None — новое своё."""
    if not row.name:
        raise RowError("Нет названия упражнения — строка пропущена.")
    if row.exercise_id is None:
        return catalog.find(row.name)
    exercise = catalog.by_pk.get(row.exercise_id)
    if exercise is None:
        # Не сопоставляем по названию молча: чужой или удалённый ID из старого
        # файла воскресил бы упражнение, которое человек сам убрал.
        raise RowError(
            f"Упражнения с ID {row.exercise_id} нет в вашем справочнике — строка пропущена. "
            "Чтобы завести его заново, очистите ID."
        )
    return exercise


def _new_fields(catalog, row):
    return {
        "name": row.name,
        "muscle_group": normalize_facet(row.muscle_group or "", catalog.groups),
        "equipment": normalize_facet(row.equipment or "", catalog.equipment),
        "measurement": row.measurement or Exercise.Measurement.WEIGHT_REPS,
    }


def _facet_change(value, current, known):
    """Новое значение группы или снаряда — или None, если менять нечего.

    Сначала сравнение как есть (с точностью до пробелов), и только потом
    приведение к принятому написанию: иначе своё «грудь» при общем «Грудь»
    переписывалось бы при каждой загрузке нетронутого файла.
    """
    if value is None or value == collapse_spaces(current):
        return None
    normalized = normalize_facet(value, known)
    return None if normalized == current else normalized


def _own_changes(catalog, exercise, row):
    """Изменения своего упражнения: только то, что в строке отличается."""
    changes = {}
    # Переименование — только по ID: без него название — ключ поиска, и другой
    # регистр значит «найди», а не «переименуй» (то же правило, что у общих).
    # Смена одного регистра по ID — законное переименование, и своё же имя
    # занятым не считается: exclude — само упражнение.
    if row.exercise_id is not None and row.name != collapse_spaces(exercise.name):
        if catalog.taken(row.name, exclude=exercise.pk):
            raise RowError(
                f"Название «{row.name}» уже занято другим упражнением — строка пропущена. "
                "Поменять имена двух упражнений местами можно в две загрузки."
            )
        changes["name"] = row.name
    group = _facet_change(row.muscle_group, exercise.muscle_group, catalog.groups)
    if group is not None:
        changes["muscle_group"] = group
    equipment = _facet_change(row.equipment, exercise.equipment, catalog.equipment)
    if equipment is not None:
        changes["equipment"] = equipment
    if row.measurement and row.measurement != exercise.measurement:
        changes["measurement"] = row.measurement
    return changes


def _differs_from_shared(catalog, exercise, row):
    """Просит ли строка поменять общее упражнение в чём-то, кроме шага.

    Название сравнивается только по ID: без него название — ключ поиска, и
    другой регистр значит не «переименуй», а «найди»."""
    if row.exercise_id is not None and row.name != collapse_spaces(exercise.name):
        return True
    if _facet_change(row.muscle_group, exercise.muscle_group, catalog.groups) is not None:
        return True
    if _facet_change(row.equipment, exercise.equipment, catalog.equipment) is not None:
        return True
    return bool(row.measurement) and row.measurement != exercise.measurement


def _step_change(catalog, report, exercise, row, measurement):
    """Новый шаг веса или None.

    Строка настроек появляется только при отличии от действующего шага:
    умолчание 2,5 ради одной загрузки не записываем. У единицы без веса шаг не
    нужен; предупреждаем, только если человек его правда поменял, — старый шаг,
    оставшийся в строке после смены единицы, шумом в отчёте не будет.
    """
    if row.weight_step is None:
        return None
    current = catalog.step(exercise) if exercise is not None else DEFAULT_WEIGHT_STEP
    if row.weight_step == current:
        return None
    if "weight_kg" not in MEASUREMENT_FIELDS[measurement]:
        report.add_warning(f"У «{row.name}» нет веса — шаг веса для него не нужен и пропущен.")
        return None
    return row.weight_step


def _summarize(report):
    """Предупреждения по итогам — одним абзацем на тему, а не строкой на упражнение."""
    if report.shared_skipped:
        report.add_warning(
            "Общие упражнения правит администратор, поэтому у них применился только шаг "
            "веса, а остальные правки пропущены: " + _names(report.shared_skipped) + "."
        )
    if report.measurement_changed:
        report.add_warning(
            "Изменена единица: "
            + _names(report.measurement_changed)
            + ". Записанные подходы остались в прежних единицах, а личный рекорд теперь "
            "считается по подходам в новой."
        )


def _names(names):
    shown = ", ".join(f"«{name}»" for name in names[:NAMES_SHOWN])
    rest = len(names) - NAMES_SHOWN
    return f"{shown} и ещё {rest}" if rest > 0 else shown
