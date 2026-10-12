"""Формат книги Excel: выгрузка истории и разбор загруженного файла.

Порядок и заголовки колонок объявлены здесь один раз (COLUMNS), и разбор берёт
их отсюда же — иначе выгрузка и загрузка разошлись бы, и файл, скачанный из
приложения, перестал бы в него загружаться.

Лист плоский: строка — подход или кардио-часть. Так проще всего заполнять руками
и переносить из чужой таблицы, а связь «тренировка → подходы» восстанавливается
группировкой по дате, времени начала и виду спорта.
"""

# Аннотации строками: поле SheetRow называется date и затенило бы одноимённый тип
# при вычислении аннотаций в момент создания класса.
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from workouts import services
from workouts.models import (
    CARDIO_ELEVATION_MAX,
    CARDIO_INCLINE_MAX,
    CARDIO_INCLINE_MIN,
    CARDIO_LEVEL_MAX,
    EXERCISE_NAME_MAX_LENGTH,
    LOCATION_NAME_MAX_LENGTH,
    MACHINE_BRAND_MAX_LENGTH,
    MACHINE_MODEL_MAX_LENGTH,
    MAX_DURATION_SEC,
    MEASUREMENT_FIELDS,
    SPORT_NAME_MAX_LENGTH,
    TIME_MEASUREMENTS,
    ExerciseMachine,
    ExerciseNote,
    StrengthSet,
    Workout,
    cardio_parts_prefetch,
    collapse_spaces,
    parse_field_value,
    rest_display,
)

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_TITLE = "Тренировки"
HELP_SHEET_TITLE = "Как заполнять"
# Та же граница, что у форм (MAX_DURATION_HOURS): тренировка в 30 часов — опечатка.
MAX_DURATION_MIN = 24 * 60


@dataclass(frozen=True)
class Column:
    key: str
    title: str
    width: int
    number_format: str | None = None


# Единственный источник правды о формате. Заголовки — то, что человек видит в
# шапке, и то, по чему разбор ищет колонки: искать по имени, а не по позиции,
# нужно затем, чтобы в файл можно было дописать свой столбец и не сломать импорт.
COLUMNS: tuple[Column, ...] = (
    Column("date", "Дата", 12, "DD.MM.YYYY"),
    Column("start", "Начало", 9, "HH:MM"),
    Column("sport", "Вид спорта", 16),
    Column("location", "Место", 18),
    # 34, а не 28: имена справочника стали точными, и самое длинное —
    # «Жим штанги на наклонной скамье» — в прежнюю ширину не помещалось.
    # Колонки под снаряд нет намеренно: он свойство справочника, а не подхода,
    # и живёт в книге справочника (exercise_excel).
    Column("exercise", "Упражнение", 34),
    Column("set_number", "Подход", 9, "0"),
    Column("weight", "Вес, кг", 10, "0.00"),
    Column("reps", "Повторы", 10, "0"),
    Column("hold", "Удержание", 12, "@"),
    Column("distance", "Дистанция, км", 15, "0.00"),
    Column("pulse", "Пульс", 9, "0"),
    Column("duration", "Длительность, мин", 18, "0"),
    Column("workout_note", "Заметка к тренировке", 30),
    Column("exercise_note", "Заметка к упражнению", 30),
    # Номер суперсета (в коде — круг, circuit) внутри тренировки: 1..k по
    # порядку. В конце и необязательная: колонки ищутся по заголовку, поэтому
    # файлы, скачанные до неё, загружаются как прежде — без суперсетов.
    Column("circuit", "Суперсет", 11, "0"),
    # Кардио-часть (CardioPart) — своя строка с пустым «Упражнением» и своими
    # видом спорта и временем: разминка на дорожке внутри силовой. «Вид спорта»
    # у неё остаётся видом спорта тренировки — по нему строки собираются в
    # тренировку, и иначе загрузка разрезала бы занятие на два. У кардио-
    # тренировки с одной частью обе колонки пусты: вид спорта и время у части
    # те же, что у тренировки, и строка выглядит как до появления частей.
    Column("cardio", "Кардио", 14),
    Column("cardio_duration", "Время кардио, мин", 18, "0"),
    # Тренажёр: у кардио-строки — тренажёр части и его настройки (решение
    # 11.10.2026), у подхода — тренажёр упражнения в месте тренировки
    # (ExerciseMachine). Тоже в конце и необязательные: старые файлы без них
    # грузятся как прежде.
    Column("machine_brand", "Производитель", 18),
    Column("machine_model", "Модель", 18),
    Column("incline", "Уклон, %", 10, "0.0"),
    Column("level", "Уровень", 9, "0"),
    # Набор высоты кардио-части — уличный двойник уклона. В конце, как все
    # поздние колонки: файлы без неё грузятся как прежде.
    Column("elevation", "Набор высоты, м", 16, "0"),
)
# Без этих трёх строка не значит ничего: дата и вид спорта опознают тренировку,
# длительность делает её записанной, а не «идущей».
REQUIRED_TITLES = ("Дата", "Вид спорта", "Длительность, мин")

HELP_LINES = (
    "Как заполнять таблицу",
    "",
    "Одна строка — один подход или одно кардио. Тренировка собирается из строк",
    "с одинаковыми датой, временем начала и видом спорта.",
    "",
    "Кардио: оставьте «Упражнение» пустым и заполните «Дистанция, км» и «Пульс».",
    "Силовая: заполните «Упражнение» и то, в чём меряется подход, — вес и повторы,",
    "либо только повторы, либо «Удержание» в виде 1:30.",
    "",
    "«Длительность, мин» нужна у каждой тренировки: без неё запись считается",
    "незаконченной, и приложение её не примет.",
    "",
    "«Подход» можно не заполнять — номера приложение проставит само по порядку строк.",
    "«Начало» тоже необязательно: пустое значит «как обычно» — полдень для прошедшего дня.",
    "",
    "Упражнения, места и виды спорта, которых в приложении ещё нет, оно заведёт само",
    "и перечислит в отчёте после загрузки.",
    "",
    "Тренировку, которая в дневнике уже есть (тот же день и вид спорта), приложение",
    "пропустит — исправленный файл можно загружать повторно без опаски.",
    "",
    "«Суперсет» — номер суперсета внутри тренировки: упражнения с одним номером",
    "выполняются по очереди, по подходу каждого. Пустой — упражнение само по себе.",
    "",
    "Кардио внутри силовой тренировки — разминку на дорожке, заминку на велосипеде —",
    "пишите отдельной строкой той же тренировки с пустым «Упражнением»: в «Кардио» —",
    "его вид спорта, в «Время кардио, мин» — сколько оно длилось. Кардио-тренировке",
    "обе колонки не нужны: вид спорта и время у неё уже есть.",
    "",
    "«Производитель» и «Модель» — тренажёр, названия как в «Моих тренажёрах»",
    "(незнакомые приложение заведёт само). У кардио это тренажёр, на котором",
    "занимались, у подхода — тренажёр упражнения в этом месте: без «Места» он не",
    "запишется, а тренажёр, уже указанный в приложении, загрузка не меняет.",
    "«Уклон, %» — от −10 до 40, «Уровень» — уровень нагрузки от 1 до 100, оба только",
    "у кардио. Всё необязательно.",
    "",
    "«Набор высоты, м» — сколько метров вверх набрано за кардио, целым числом.",
    "Необязательно; ноль — то же, что пусто.",
    "",
    "Упражнение с повторами без веса приложение заводит как «вес × повторы»: вес 0",
    "у него значит «со своим весом». Если это чистые повторы — поменяйте единицу",
    "на странице упражнения.",
)

HEAD_FILL = PatternFill("solid", fgColor="EFEFEF")


def _text(cell, value):
    """Записать текст текстом.

    openpyxl превращает строку, начинающуюся с «=», в формулу. Заметка «=2+2»
    должна остаться заметкой, а не считаться при открытии файла.
    """
    cell.value = value or ""
    cell.data_type = "s"


def sheet_rows(user):
    """Строки листа по всей записанной истории пользователя.

    Пять запросов и ни одного в цикле: тренировки с кардио-частями, подходы,
    заметки и тренажёры выбираются пачкой на всю историю и раскладываются по
    тренировкам в Python. Порядок упражнений внутри тренировки — общий с
    экраном итога (services.group_sets), поэтому выгрузка не изобретает
    собственного.
    """
    workouts = list(
        Workout.objects.filter(user=user)
        .finished()
        .select_related("sport", "location")
        .prefetch_related(cardio_parts_prefetch())
        .order_by("started_at", "id")
    )
    if not workouts:
        return []

    sets_by_workout = {}
    for row in (
        StrengthSet.objects.filter(workout__user=user, workout__duration_min__isnull=False)
        .select_related("exercise")
        .order_by("workout_id", "id")
    ):
        sets_by_workout.setdefault(row.workout_id, []).append(row)

    notes_by_workout = {}
    for workout_id, exercise_id, text in ExerciseNote.objects.filter(
        workout__user=user
    ).values_list("workout_id", "exercise_id", "text"):
        notes_by_workout.setdefault(workout_id, {})[exercise_id] = text

    # Тренажёры упражнений по местам: словарём, а не аннотацией подходов — место
    # у каждой тренировки своё, и подзапрос пришлось бы считать на каждый подход.
    machines = {
        (machine.exercise_id, machine.location_id): (
            machine.brand.name,
            machine.model.name if machine.model_id else "",
        )
        for machine in ExerciseMachine.objects.filter(user=user).select_related("brand", "model")
    }

    rows = []
    for workout in workouts:
        rows.extend(
            _workout_rows(
                workout,
                sets_by_workout.get(workout.pk, []),
                notes_by_workout.get(workout.pk, {}),
                machines,
            )
        )
    return rows


def _workout_rows(workout, sets, notes, machines):
    """Строки одной тренировки: по строке на подход и на кардио-часть.

    Порядок тот же, что на итоге: упражнения, за ними кардио в порядке ввода.
    Тренировка без того и другого («плавал 40 минут, не мерил») — одна строка.
    """
    started = timezone.localtime(workout.started_at)
    common = {
        "date": started.date(),
        "start": started.time().replace(second=0, microsecond=0),
        "sport": workout.sport.name,
        "location": workout.location.name if workout.location else "",
        # Длительность повторяется в каждой строке: человеку, который смотрит на
        # середину тренировки, она нужна там же, а не в первой строке сверху.
        "duration": workout.duration_min,
    }
    rows = []
    for group in services.group_sets(sets, notes):
        # Тренажёр — тот, что указан у упражнения в месте тренировки сейчас:
        # ExerciseMachine — настройка, а не снимок, и итог старой тренировки
        # показывает его же. Пишется в каждую строку, как место: по нему
        # фильтруют, а повторы короткого названия стеной не выглядят.
        brand, model = machines.get((group["exercise"].pk, workout.location_id), ("", ""))
        for index, item in enumerate(group["sets"]):
            row = dict(common)
            row["exercise"] = group["exercise"].name
            row["set_number"] = item.display_number
            fields = MEASUREMENT_FIELDS[item.measurement]
            if "weight_kg" in fields:
                row["weight"] = item.weight_kg
            if "reps" in fields:
                row["reps"] = item.reps
            if item.measurement in TIME_MEASUREMENTS:
                # Удержание пишем как «1:30» — тем же форматом, каким его видно в
                # приложении. Число секунд в таблице читается хуже и в чужой
                # таблице почти не встречается.
                row["hold"] = rest_display(item.duration_sec)
            # Заметки — только в первой строке своей сущности: в каждой строке они
            # превратили бы лист в стену повторов.
            if index == 0:
                row["exercise_note"] = group["note"]
            row["circuit"] = group["circuit_no"]
            row["machine_brand"] = brand
            row["machine_model"] = model
            rows.append(row)
    parts = workout.cardio_parts.all()
    # Единственная часть тренировки без подходов — это сама тренировка (так её
    # записывает форма кардио), и строка у неё та же, что до появления частей.
    lone = not sets and len(parts) == 1
    rows.extend(_cardio_row(common, workout, part, lone=lone) for part in parts)
    if not rows:
        rows.append(dict(common))
    rows[0]["workout_note"] = workout.note
    return rows


def _cardio_row(common, workout, part, *, lone):
    """Строка кардио-части. Вид спорта и время пишутся, только когда они свои:
    пустые загрузка читает как «те же, что у тренировки»."""
    row = dict(common)
    row["distance"] = part.distance_km
    row["pulse"] = part.avg_heart_rate
    row["machine_brand"] = part.brand.name if part.brand_id else ""
    row["machine_model"] = part.model.name if part.model_id else ""
    row["incline"] = part.incline_pct
    row["level"] = part.level
    row["elevation"] = part.elevation_gain_m
    if part.sport_id != workout.sport_id:
        row["cardio"] = part.sport.name
    if not (lone and part.duration_min == workout.duration_min):
        row["cardio_duration"] = part.duration_min
    return row


def write_table(sheet, columns, rows):
    """Лист-таблица: шапка из колонок и строки словарями по их ключам.

    Общий для обеих книг — истории и справочника упражнений: оформление шапки и
    правило «текст пишется текстом» должны быть одинаковыми, а не двумя копиями.
    Колонка без формата (или «@») пишется текстом, с форматом — числом.
    """
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        _text(cell, column.title)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = column.width

    for number, values in enumerate(rows, start=2):
        for index, column in enumerate(columns, start=1):
            value = values.get(column.key)
            if value is None or value == "":
                continue
            cell = sheet.cell(row=number, column=index)
            if column.number_format in (None, "@"):
                _text(cell, value)
            else:
                cell.value = value
            if column.number_format:
                cell.number_format = column.number_format

    # Шапка закреплена и с автофильтром: в истории за год строк тысячи, и без
    # этого файл неудобен уже на второй минуте.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 1)}"


def write_help(book, lines):
    """Лист «Как заполнять»: первая строка — заголовок, дальше абзацы по строке."""
    help_sheet = book.create_sheet(HELP_SHEET_TITLE)
    help_sheet.column_dimensions["A"].width = 90
    for number, line in enumerate(lines, start=1):
        cell = help_sheet.cell(row=number, column=1)
        _text(cell, line)
        if number == 1:
            cell.font = Font(bold=True)


def build_workbook(user):
    """Книга целиком: лист с историей и лист с подсказкой."""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE
    write_table(sheet, COLUMNS, sheet_rows(user))
    write_help(book, HELP_LINES)
    return book


def export_filename(today):
    """Имя файла: «Дневник тренировок 2026-09-25.xlsx».

    Дата в имени, чтобы две выгрузки в папке «Загрузки» не путались. Кириллица
    допустима: FileResponse отдаёт имя в кодировке RFC 5987.
    """
    return f"Дневник тренировок {today:%Y-%m-%d}.xlsx"


# --- Разбор загруженной книги -------------------------------------------------

MAX_ROWS = 20_000


class WorkbookError(Exception):
    """Книгу нельзя прочитать целиком; args[0] — текст для человека."""


# Сколько ошибок и предупреждений показываем. Остальные только считаются: список
# на тысячу строк никто не прочтёт, а страница с ним грузилась бы вечность.
ERRORS_KEPT = 200


@dataclass
class ReportLog:
    """Журнал загрузки: ошибки по строкам и предупреждения.

    Общий для отчётов обеих книг — истории и справочника: «Строка N: …» и
    предел показанного должны выглядеть одинаково, где бы ни случились.
    """

    error_count: int = 0
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def hidden_errors(self):
        return max(0, self.error_count - len(self.errors))

    def add_error(self, number, message):
        self.error_count += 1
        if len(self.errors) < ERRORS_KEPT:
            self.errors.append(f"Строка {number}: {message}")

    def add_warning(self, message):
        if message not in self.warnings and len(self.warnings) < ERRORS_KEPT:
            self.warnings.append(message)


@dataclass
class SheetRow:
    """Строка листа, уже приведённая к типам модели.

    Имена полей совпадают с полями StrengthSet (weight_kg, reps, duration_sec):
    так строка годится как источник для services.set_values без переходника.
    Ошибки разбора копятся в errors, а не бросаются: из-за одной опечатки в
    трёхсотстрочной таблице терять остальное незачем.
    """

    number: int
    date: date | None = None
    start_time: time | None = None
    sport: str = ""
    location: str = ""
    exercise: str = ""
    weight_kg: Decimal = Decimal(0)
    reps: int = 0
    duration_sec: int = 0
    distance_km: Decimal | None = None
    avg_heart_rate: int | None = None
    duration_min: int | None = None
    workout_note: str = ""
    exercise_note: str = ""
    circuit: int | None = None
    cardio_sport: str = ""
    cardio_duration_min: int | None = None
    machine_brand: str = ""
    machine_model: str = ""
    incline_pct: Decimal | None = None
    level: int | None = None
    elevation_gain_m: int | None = None
    errors: list[str] = field(default_factory=list)
    # Строка, которую не к чему отнести: название упражнения или вида спорта не
    # помещается в справочник. Такая не загружается вовсе — с пустым названием
    # подход стал бы кардио-строкой, а тренировка осталась бы без вида спорта.
    broken: bool = False

    @property
    def is_cardio(self):
        """Пустое упражнение и есть признак кардио-строки."""
        return not self.exercise

    @property
    def has_cardio_part(self):
        """Описывает ли кардио-строка часть: заполнено хоть что-то о ней.

        Пустая — кардио-тренировка без подробностей («плавал 40 минут, не
        мерил»): части у неё нет, как и у формы кардио без дистанции. Ноль
        дистанции — тоже «не мерил», а нулевой уклон — настоящий уклон.
        """
        return bool(
            self.cardio_sport
            or self.cardio_duration_min
            or self.distance_km
            or self.avg_heart_rate
            or self.machine_brand
            or self.machine_model
            or self.incline_pct is not None
            or self.level
            or self.elevation_gain_m
        )


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return collapse_spaces(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value).strip()


def cell_limited(value, *, limit, what):
    """Текст ячейки не длиннее поля в базе.

    Общий для обеих книг. Обрезать молча нельзя — человек не узнает, а
    пропущенное без проверки длинное название уронило бы загрузку ошибкой базы
    вместо строки в отчёте.
    """
    text = cell_text(value)
    if len(text) > limit:
        raise ValueError(f"{what} длиннее {limit} символов.")
    return text


def _cell_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError("Не похоже на дату — напишите 25.09.2026.")


def _cell_time(value):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, timedelta):
        total = int(value.total_seconds()) % 86400
        return time(total // 3600, total % 3600 // 60)
    text = cell_text(value)
    if not text:
        return None
    for pattern in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).time().replace(second=0)
        except ValueError:
            continue
    raise ValueError("Не похоже на время — напишите 18:30.")


def _cell_seconds(value):
    """Удержание в секунды. «1:30» — это полторы минуты, как на экране.

    Excel сам решает, чем считать введённое: «1:30» он часто превращает во
    время 01:30:00. Читаем такие значения как «минуты:секунды» — час удержания
    и так за границей MAX_DURATION_SEC, а полторы минуты в планке обычное дело.
    """
    if value is None or value == "":
        return 0
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second:
            return min(MAX_DURATION_SEC, value.hour * 3600 + value.minute * 60 + value.second)
        return min(MAX_DURATION_SEC, value.hour * 60 + value.minute)
    if isinstance(value, timedelta):
        return min(MAX_DURATION_SEC, int(value.total_seconds()))
    if isinstance(value, (int, float, Decimal)):
        return min(MAX_DURATION_SEC, max(0, int(value)))
    return parse_field_value("duration_sec", cell_text(value))


def _cell_weight(value):
    if value is None or value == "":
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)):
        return parse_field_value("weight_kg", str(value))
    return parse_field_value("weight_kg", cell_text(value))


def _cell_int(value, *, what):
    if value is None or value == "":
        return None
    # «inf» и 1e999 дают OverflowError, «nan» — ValueError с английским текстом;
    # и то и другое должно стать человеческой ошибкой строки, а не пятисоткой.
    try:
        if isinstance(value, (int, float, Decimal)):
            return int(value)
        return int(float(cell_text(value).replace(",", ".")))
    except (ValueError, OverflowError) as error:
        raise ValueError(f"{what} — это целое число.") from error


def _cell_circuit(value):
    """Номер суперсета: пусто — упражнение само по себе, иначе 1..99."""
    number = _cell_int(value, what="Суперсет")
    if number is not None and not 1 <= number <= 99:
        raise ValueError("Суперсет — номер от 1 до 99.")
    return number


def _cell_incline(value):
    """Уклон кардио-тренажёра: −10…40 %, один знак после запятой."""
    number = _cell_decimal(value, what="Уклон")
    if number is None:
        return None
    if not CARDIO_INCLINE_MIN <= number <= CARDIO_INCLINE_MAX:
        raise ValueError("Уклон — от −10 до 40 %.")
    return number.quantize(Decimal("0.1"))


def _cell_level(value):
    """Уровень нагрузки кардио-тренажёра: 1…CARDIO_LEVEL_MAX."""
    number = _cell_int(value, what="Уровень")
    if number is not None and not 1 <= number <= CARDIO_LEVEL_MAX:
        raise ValueError(f"Уровень — от 1 до {CARDIO_LEVEL_MAX}.")
    return number


def _cell_elevation(value):
    """Набор высоты кардио-части: целые метры. Пусто и ноль — «не указан», как в форме."""
    number = _cell_int(value, what="Набор высоты")
    if not number:
        return None
    if not 0 < number <= CARDIO_ELEVATION_MAX:
        raise ValueError(f"Набор высоты — от 1 до {CARDIO_ELEVATION_MAX} м.")
    return number


def _cell_cardio_minutes(value):
    """Время кардио-части в минутах. Пусто и ноль — «не указано», как в форме."""
    number = _cell_int(value, what="Время кардио")
    if not number:
        return None
    if not 0 < number <= MAX_DURATION_MIN:
        raise ValueError(f"Время кардио — от 1 до {MAX_DURATION_MIN} минут.")
    return number


def _cell_decimal(value, *, what):
    if value is None or value == "":
        return None
    text = cell_text(value).replace(",", ".") if not isinstance(value, str) else value
    try:
        number = Decimal(str(text).replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"{what} — это число, например 7,2.") from error
    # «nan» Decimal округляет без возражений, а numeric в Postgres его примет —
    # в дистанции тренировки оказалось бы «не число».
    if not number.is_finite():
        raise ValueError(f"{what} — это число, например 7,2.")
    return number


def header_index(titles, *, columns, required_titles):
    """Колонки ищутся по заголовку, а не по позиции.

    Так в файл можно дописать свой столбец или переставить их местами, и он
    всё равно загрузится: человек, ведущий таблицу годами, скорее переименует
    шапку, чем подгонит порядок под наш. Колонки, которой в файле нет, в индексе
    нет тоже (None) — разбор строки решает, что это значит.
    """
    index = {}
    for position, title in enumerate(titles):
        name = cell_text(title).lower()
        if name and name not in index:
            index[name] = position
    missing = [title for title in required_titles if title.lower() not in index]
    if missing:
        raise WorkbookError("В первой строке не хватает колонок: " + ", ".join(missing) + ".")
    return {column.key: index.get(column.title.lower()) for column in columns}


def cell_getter(values, index):
    """Ячейка строки по ключу колонки; None — если колонки нет или строка короче."""

    def cell(key):
        position = index.get(key)
        if position is None or position >= len(values):
            return None
        return values[position]

    return cell


NOT_A_WORKBOOK = (
    "Не получилось открыть файл. Нужен .xlsx — сохраните таблицу из Excel "
    "как «Книга Excel (.xlsx)»."
)


def read_rows(
    stream,
    *,
    sheet_title,
    columns,
    required_titles,
    parse_row,
    max_rows,
    check_titles=None,
    optional=False,
):
    """Строки книги без записи в базу. Кидает WorkbookError, если читать нечего.

    Общий для обеих книг: открыть из памяти, найти лист и шапку, пропустить пустые
    строки, соблюсти лимит. Что значит строка, решает parse_row(number, values,
    index) — номер в нём настоящий, как его видит человек в Excel. check_titles —
    проверка шапки сверх обязательных колонок: так справочник узнаёт по шапке
    файл с историей, который его обязательную колонку тоже содержит.
    optional — дополнительный лист книги: нет листа с таким названием — нет и
    строк (без подмены первым листом, как у основного).
    """
    try:
        book = load_workbook(stream, read_only=True, data_only=True)
    except Exception as error:
        # Битый zip, .xls, csv, картинка — исключений тут зоопарк, а ответ один:
        # это не та книга. Падать пятисоткой на кривом файле нельзя.
        raise WorkbookError(NOT_A_WORKBOOK) from error
    try:
        if optional and sheet_title not in book.sheetnames:
            return []
        try:
            sheet = book[sheet_title] if sheet_title in book.sheetnames else book.worksheets[0]
        except IndexError as error:
            raise WorkbookError(NOT_A_WORKBOOK) from error
        lines = _safe_lines(sheet.iter_rows(values_only=True))
        try:
            titles = next(lines)
        except StopIteration as error:
            raise WorkbookError("Файл пустой: в нём нет даже шапки.") from error
        if check_titles is not None:
            check_titles(titles)
        index = header_index(titles, columns=columns, required_titles=required_titles)
        rows = []
        for number, values in enumerate(lines, start=2):
            if all(value is None or value == "" for value in values):
                continue
            if len(rows) >= max_rows:
                raise WorkbookError(
                    f"Слишком много строк: больше {max_rows} за раз не примем. "
                    "Разделите таблицу на части."
                )
            rows.append(parse_row(number, values, index))
        return rows
    finally:
        book.close()


def _safe_lines(lines):
    """Строки листа, где сбой чтения — это WorkbookError, а не пятисотка.

    В режиме read_only openpyxl разбирает лист лениво: битый XML внутри книги
    всплывает не при открытии, а на очередной строке. Ловим только чтение —
    ошибки разбора строки в parse_row сюда не попадают и не маскируются.
    """
    while True:
        try:
            values = next(lines)
        except StopIteration:
            return
        except Exception as error:
            raise WorkbookError(NOT_A_WORKBOOK) from error
        yield values


def read_sheet(stream, *, max_rows=None):
    """Строки книги с историей. Кидает WorkbookError, если читать нечего."""
    # Лимит разрешается при вызове, а не в сигнатуре: иначе значение защёлкнулось
    # бы на импорте модуля, и подменить его (в тестах или настройкой) было бы нечем.
    max_rows = MAX_ROWS if max_rows is None else max_rows
    return read_rows(
        stream,
        sheet_title=SHEET_TITLE,
        columns=COLUMNS,
        required_titles=REQUIRED_TITLES,
        parse_row=_sheet_row,
        max_rows=max_rows,
    )


def _sheet_row(number, values, index):
    row = SheetRow(number=number)
    cell = cell_getter(values, index)

    def read(key, parser, target, *, vital=False):
        try:
            setattr(row, target, parser(cell(key)))
        except ValueError as error:
            row.errors.append(str(error))
            row.broken = row.broken or vital

    def name(limit, what):
        return lambda value: cell_limited(value, limit=limit, what=what)

    read("date", _cell_date, "date")
    read("start", _cell_time, "start_time")
    # Вид спорта и упражнение решают, куда пойдёт строка, — поэтому vital.
    read("sport", name(SPORT_NAME_MAX_LENGTH, "Вид спорта"), "sport", vital=True)
    read("location", name(LOCATION_NAME_MAX_LENGTH, "Место"), "location")
    read("exercise", name(EXERCISE_NAME_MAX_LENGTH, "Упражнение"), "exercise", vital=True)
    read("weight", _cell_weight, "weight_kg")
    read("reps", lambda value: _cell_int(value, what="Повторы") or 0, "reps")
    read("hold", _cell_seconds, "duration_sec")
    read("distance", lambda value: _cell_decimal(value, what="Дистанция"), "distance_km")
    read("pulse", lambda value: _cell_int(value, what="Пульс"), "avg_heart_rate")
    read("duration", lambda value: _cell_int(value, what="Длительность"), "duration_min")
    row.workout_note = cell_text(cell("workout_note"))
    row.exercise_note = cell_text(cell("exercise_note"))
    read("circuit", _cell_circuit, "circuit")
    read("cardio", name(SPORT_NAME_MAX_LENGTH, "Кардио"), "cardio_sport")
    read("cardio_duration", _cell_cardio_minutes, "cardio_duration_min")
    read("machine_brand", name(MACHINE_BRAND_MAX_LENGTH, "Производитель"), "machine_brand")
    read("machine_model", name(MACHINE_MODEL_MAX_LENGTH, "Модель"), "machine_model")
    read("incline", _cell_incline, "incline_pct")
    read("level", _cell_level, "level")
    read("elevation", _cell_elevation, "elevation_gain_m")
    return row
